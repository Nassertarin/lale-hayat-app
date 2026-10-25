# -*- coding: utf-8 -*-
"""
تولید PDF و Excel فرم گزارش روزانه از دیتابیس (بر اساس form_date)
"""

import os
import csv
from datetime import datetime
from fpdf import FPDF

import arabic_reshaper
from bidi.algorithm import get_display

try:
    import jdatetime
    HAS_JDATETIME = True
except ImportError:
    HAS_JDATETIME = False


BASE_DIR = os.path.dirname(os.path.abspath(__file__))

FONT_REG_PATH = os.path.join(BASE_DIR, "assets", "fonts", "BNazanin.ttf")
FONT_BOLD_PATH = os.path.join(BASE_DIR, "assets", "fonts", "BNazaninBold.ttf")
LOGO_PATH = os.path.join(BASE_DIR, "assets", "images", "logo.png")

REPORTS_DIR = os.path.join(BASE_DIR, "reports")


def fa(text):
    if not text:
        return ""
    reshaped = arabic_reshaper.reshape(str(text))
    return get_display(reshaped)


def safe_float(value, default=0.0):
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


BLACK = (0, 0, 0)
RED = (200, 0, 0)
GREEN = (0, 100, 0)
WHITE = (255, 255, 255)

LEFT_EDGE = 12
RIGHT_EDGE = 198
TOTAL_WIDTH = RIGHT_EDGE - LEFT_EDGE   # 186mm

MORGH_TABLE_WIDTH = TOTAL_WIDTH
MORGH_LEFT = RIGHT_EDGE - MORGH_TABLE_WIDTH


# ==========================================
# کلاس PDF
# ==========================================

class DailyFormPDF(FPDF):

    def __init__(self):
        super().__init__(orientation="P", unit="mm", format="A4")
        self.set_auto_page_break(auto=False)
        self.set_margins(12, 12, 12)

        if not os.path.exists(FONT_REG_PATH):
            raise FileNotFoundError(f"فونت پیدا نشد: {FONT_REG_PATH}")

        self.add_font("Nazanin", "", FONT_REG_PATH)

        if os.path.exists(FONT_BOLD_PATH):
            self.add_font("Nazanin", "B", FONT_BOLD_PATH)
            self.has_bold = True
        else:
            self.has_bold = False

    def draw_header(self, form_date, day_name=""):
        if os.path.exists(LOGO_PATH):
            try:
                self.image(LOGO_PATH, x=14, y=8, w=18)
            except Exception:
                pass

        self.set_y(10)
        if self.has_bold:
            self.set_font("Nazanin", "B", size=17)
        else:
            self.set_font("Nazanin", size=17)

        self.set_text_color(*BLACK)
        self.cell(186, 11, fa("فرم گزارش روزانه کشتارگاه"), align="C")

        self.set_y(28)
        self.set_font("Nazanin", size=11)
        self.set_text_color(*BLACK)

        date_text = f"تاریخ: {form_date}"
        if day_name:
            date_text += f"  ({day_name})"

        self.set_xy(120, 28)
        self.cell(78, 6, fa(date_text), align="R")

        self.set_y(38)

    def draw_section_title(self, title, y):
        if self.has_bold:
            self.set_font("Nazanin", "B", size=12)
        else:
            self.set_font("Nazanin", size=12)

        self.set_text_color(*BLACK)
        self.set_xy(LEFT_EDGE, y)
        self.cell(TOTAL_WIDTH, 7, fa(title), align="R")

    def draw_packed_table(self, rows, y):
        self.draw_section_title("مرغ پرکنده", y)
        y += 8

        unit = TOTAL_WIDTH / 15

        widths = [
            3.5 * unit,
            4.5 * unit,
            5.0 * unit,
            2.0 * unit,
        ]
        headers = [
            "وزن کل",
            "شماره گاری",
            "نام خریدار",
            "ردیف",
        ]

        self.set_font("Nazanin", size=9)
        self.set_text_color(*BLACK)
        self.set_draw_color(*BLACK)

        self.set_xy(LEFT_EDGE, y)
        for h, w in zip(headers, widths):
            self.cell(w, 8.5, fa(h), border=1, align="C")
        self.ln()

        self.set_font("Nazanin", size=9)

        total_weight = 0.0

        for row in rows:
            cart_numbers = ", ".join(
                c["cart_number"] for c in row.get("carts", [])
            )
            weight = safe_float(row.get("total_net_weight", 0))

            self.cell(widths[0], 7, f"{weight:,.2f}",
                      border=1, align="C")
            self.cell(widths[1], 7, fa(cart_numbers),
                      border=1, align="C")
            self.cell(widths[2], 7, fa(row.get("buyer_name", "")),
                      border=1, align="R")
            self.cell(widths[3], 7, str(row.get("row_num", "")),
                      border=1, align="C")
            self.ln()

            total_weight += weight

        self.set_font("Nazanin", size=9)

        self.cell(widths[0], 7.5, f"{total_weight:,.2f}",
                  border=1, align="C")
        self.cell(widths[1] + widths[2] + widths[3], 7.5,
                  fa("جمع"), border=1, align="C")

        self.ln(10)
        return self.get_y()

    def draw_alayeshat_table(self, rows, y):
        self.draw_section_title("آلایشات", y)
        y += 8

        unit = TOTAL_WIDTH / 14

        widths = [
            2.0 * unit,
            2.0 * unit,
            2.0 * unit,
            2.0 * unit,
            2.0 * unit,
            3.0 * unit,
            1.0 * unit,
        ]
        headers = [
            "کله", "پا", "سنگدان", "جگر", "دل",
            "نام مشتری", "ردیف",
        ]

        self.set_font("Nazanin", size=9)
        self.set_text_color(*BLACK)
        self.set_draw_color(*BLACK)

        self.set_xy(LEFT_EDGE, y)
        for h, w in zip(headers, widths):
            self.cell(w, 8.5, fa(h), border=1, align="C")
        self.ln()

        self.set_font("Nazanin", size=9)

        total_kale = 0.0
        total_pa = 0.0
        total_sangdane = 0.0
        total_jigar = 0.0
        total_del = 0.0

        for row in rows:
            kale_f = safe_float(row.get("kale_weight", 0))
            pa_f = safe_float(row.get("pa_weight", 0))
            sangdane_f = safe_float(row.get("sangdane_weight", 0))
            jigar_f = safe_float(row.get("jigar_weight", 0))
            del_f = safe_float(row.get("del_weight", 0))

            self.cell(widths[0], 7,
                      f"{kale_f:,.2f}" if kale_f else "",
                      border=1, align="C")
            self.cell(widths[1], 7,
                      f"{pa_f:,.2f}" if pa_f else "",
                      border=1, align="C")
            self.cell(widths[2], 7,
                      f"{sangdane_f:,.2f}" if sangdane_f else "",
                      border=1, align="C")
            self.cell(widths[3], 7,
                      f"{jigar_f:,.2f}" if jigar_f else "",
                      border=1, align="C")
            self.cell(widths[4], 7,
                      f"{del_f:,.2f}" if del_f else "",
                      border=1, align="C")
            self.cell(widths[5], 7,
                      fa(row.get("customer_name", "")),
                      border=1, align="R")
            self.cell(widths[6], 7, str(row.get("row_num", "")),
                      border=1, align="C")
            self.ln()

            total_kale += kale_f
            total_pa += pa_f
            total_sangdane += sangdane_f
            total_jigar += jigar_f
            total_del += del_f

        self.set_font("Nazanin", size=9)

        self.cell(widths[0], 7.5, f"{total_kale:,.2f}",
                  border=1, align="C")
        self.cell(widths[1], 7.5, f"{total_pa:,.2f}",
                  border=1, align="C")
        self.cell(widths[2], 7.5, f"{total_sangdane:,.2f}",
                  border=1, align="C")
        self.cell(widths[3], 7.5, f"{total_jigar:,.2f}",
                  border=1, align="C")
        self.cell(widths[4], 7.5, f"{total_del:,.2f}",
                  border=1, align="C")
        self.cell(widths[5], 7.5, "",
                  border=1, align="C")
        self.cell(widths[6], 7.5, fa("جمع"),
                  border=1, align="C")

        self.ln(10)
        return self.get_y(), {
            "kale": total_kale,
            "pa": total_pa,
            "sangdane": total_sangdane,
            "jigar": total_jigar,
            "del": total_del,
        }

    def draw_morgh_table(self, rows, y):
        self.draw_section_title("مرغ اصلاحی", y)
        y += 8

        unit = MORGH_TABLE_WIDTH / 20

        widths = [
            2.4 * unit,
            2.6 * unit,
            1.7 * unit,
            1.7 * unit,
            1.7 * unit,
            1.7 * unit,
            1.9 * unit,
            1.9 * unit,
            3.0 * unit,
            1.4 * unit,
        ]

        headers = [
            "ضایعات", "مرغ کامل بسته‌بندی", "گردن", "بال",
            "ران", "سینه", "قطعه", "مرغ لکه",
            "نام مشتری", "ردیف",
        ]

        self.set_font("Nazanin", size=8)
        self.set_text_color(*BLACK)
        self.set_draw_color(*BLACK)

        self.set_xy(MORGH_LEFT, y)
        for h, w in zip(headers, widths):
            self.cell(w, 9.5, fa(h), border=1, align="C")
        self.ln()

        self.set_font("Nazanin", size=8)

        total_waste = 0.0
        total_whole = 0.0
        total_gardan = 0.0
        total_bal = 0.0
        total_ran = 0.0
        total_sine = 0.0
        total_piece = 0.0
        total_lake = 0.0

        for row in rows:
            waste_f = safe_float(row.get("waste_weight", 0))
            whole_f = safe_float(row.get("whole_weight", 0))
            gardan_f = safe_float(row.get("gardan_weight", 0))
            bal_f = safe_float(row.get("bal_weight", 0))
            ran_f = safe_float(row.get("ran_weight", 0))
            sine_f = safe_float(row.get("sine_weight", 0))
            piece_f = safe_float(row.get("piece_weight", 0))
            lake_f = safe_float(row.get("lake_weight", 0))

            self.set_x(MORGH_LEFT)

            self.cell(widths[0], 7, f"{waste_f:,.2f}",
                      border=1, align="C")
            self.cell(widths[1], 7, f"{whole_f:,.2f}",
                      border=1, align="C")
            self.cell(widths[2], 7, f"{gardan_f:,.2f}",
                      border=1, align="C")
            self.cell(widths[3], 7, f"{bal_f:,.2f}",
                      border=1, align="C")
            self.cell(widths[4], 7, f"{ran_f:,.2f}",
                      border=1, align="C")
            self.cell(widths[5], 7, f"{sine_f:,.2f}",
                      border=1, align="C")
            self.cell(widths[6], 7, f"{piece_f:,.2f}",
                      border=1, align="C")
            self.cell(widths[7], 7, f"{lake_f:,.2f}",
                      border=1, align="C")
            self.cell(widths[8], 7, fa(row.get("customer_name", "")),
                      border=1, align="R")
            self.cell(widths[9], 7, str(row.get("row_num", "")),
                      border=1, align="C")
            self.ln()

            total_waste += waste_f
            total_whole += whole_f
            total_gardan += gardan_f
            total_bal += bal_f
            total_ran += ran_f
            total_sine += sine_f
            total_piece += piece_f
            total_lake += lake_f

        self.set_font("Nazanin", size=8)
        self.set_x(MORGH_LEFT)

        self.cell(widths[0], 7.5, f"{total_waste:,.2f}",
                  border=1, align="C")
        self.cell(widths[1], 7.5, f"{total_whole:,.2f}",
                  border=1, align="C")
        self.cell(widths[2], 7.5, f"{total_gardan:,.2f}",
                  border=1, align="C")
        self.cell(widths[3], 7.5, f"{total_bal:,.2f}",
                  border=1, align="C")
        self.cell(widths[4], 7.5, f"{total_ran:,.2f}",
                  border=1, align="C")
        self.cell(widths[5], 7.5, f"{total_sine:,.2f}",
                  border=1, align="C")
        self.cell(widths[6], 7.5, f"{total_piece:,.2f}",
                  border=1, align="C")
        self.cell(widths[7], 7.5, f"{total_lake:,.2f}",
                  border=1, align="C")
        self.cell(widths[8], 7.5, "",
                  border=1, align="C")
        self.cell(widths[9], 7.5, fa("جمع"),
                  border=1, align="C")

        self.ln(10)

        return self.get_y(), {
            "waste": total_waste,
            "whole": total_whole,
            "piece": total_piece,
            "lake": total_lake,
        }

    def draw_calc_and_discrepancies(self, alayeshat_totals, morgh_totals, waste, y):
        self.draw_section_title("محاسبه و مغایرت‌ها", y)
        y += 9

        self.set_font("Nazanin", size=10)
        self.set_text_color(*BLACK)

        self.set_xy(LEFT_EDGE, y)
        self.cell(
            TOTAL_WIDTH, 6.5,
            fa("(جمع مرغ کامل + جمع قطعه) - جمع مرغ لکه = ضایعات"),
            align="R",
        )
        y += 6.5

        self.set_xy(LEFT_EDGE, y)
        self.cell(
            TOTAL_WIDTH, 6.5,
            fa(f"({morgh_totals['whole']:,.2f} + {morgh_totals['piece']:,.2f}) - {morgh_totals['lake']:,.2f} = {waste:,.2f}"),
            align="R",
        )
        y += 6.5

        if self.has_bold:
            self.set_font("Nazanin", "B", size=11)
        else:
            self.set_font("Nazanin", size=11)

        self.set_xy(LEFT_EDGE, y)
        self.cell(
            TOTAL_WIDTH, 7,
            fa(f"{waste:,.2f} کیلوگرم = ضایعات"),
            align="R",
        )
        y += 10

        total_assigned = (
            alayeshat_totals["del"] +
            alayeshat_totals["jigar"] +
            alayeshat_totals["sangdane"] +
            alayeshat_totals["pa"] +
            alayeshat_totals["kale"]
        )

        total_weighing = getattr(self, "total_weighing_alayeshat", 0.0)
        diff = total_assigned - total_weighing

        if self.has_bold:
            self.set_font("Nazanin", "B", size=11)
        else:
            self.set_font("Nazanin", size=11)

        self.set_xy(LEFT_EDGE, y)
        if abs(diff) > 0.01:
            self.set_text_color(*RED)
            self.cell(
                TOTAL_WIDTH, 7,
                fa(f"مغایرت آلایشات: جمع تخصیص {total_assigned:,.2f} - جمع توزین {total_weighing:,.2f} = {diff:+,.2f} کیلوگرم"),
                align="R",
            )
        else:
            self.set_text_color(*GREEN)
            self.cell(
                TOTAL_WIDTH, 7,
                fa(f"مغایرت آلایشات: بدون مغایرت (جمع {total_assigned:,.2f} کیلوگرم)"),
                align="R",
            )
        y += 10

        piece_diffs = getattr(self, "piece_diffs", [])

        self.set_text_color(*BLACK)
        self.set_xy(LEFT_EDGE, y)

        if piece_diffs:
            self.set_text_color(*RED)
            self.cell(
                TOTAL_WIDTH, 7,
                fa(f"مغایرت قطعه: {len(piece_diffs)} ردیف دارای اختلاف"),
                align="R",
            )
            y += 7
            for line in piece_diffs:
                self.set_xy(LEFT_EDGE, y)
                self.set_font("Nazanin", size=9)
                self.cell(TOTAL_WIDTH, 6, fa(line), align="R")
                y += 6
        else:
            self.set_text_color(*GREEN)
            self.cell(
                TOTAL_WIDTH, 7,
                fa("مغایرت قطعه: بدون مغایرت"),
                align="R",
            )
            y += 7

        return y

    # ── تابع جدید: مغایرت گاری‌ها ──
    def draw_cart_discrepancies(self, cart_diffs, y):
        """
        نمایش مغایرت گاری‌ها زیر مغایرت قطعه.
        اگر cart_diffs خالی باشد → «بدون مغایرت» (سبز)
        در غیر این صورت → پیام قرمز + لیست مغایرت‌ها
        """
        self.set_text_color(*BLACK)

        if cart_diffs:
            self.set_text_color(*RED)
            if self.has_bold:
                self.set_font("Nazanin", "B", size=11)
            else:
                self.set_font("Nazanin", size=11)

            self.set_xy(LEFT_EDGE, y)
            self.cell(
                TOTAL_WIDTH, 7,
                fa(f"مغایرت گاری‌ها: {len(cart_diffs)} مورد"),
                align="R",
            )
            y += 7

            self.set_font("Nazanin", size=9)
            self.set_text_color(*RED)

            for line in cart_diffs:
                # شکستن خطوط طولانی
                if len(line) > 95:
                    words = line.split(" ")
                    current = ""
                    for word in words:
                        if len(current) + len(word) + 1 > 95:
                            self.set_xy(LEFT_EDGE, y)
                            self.cell(TOTAL_WIDTH, 6, fa(current), align="R")
                            y += 6
                            current = word + " "
                        else:
                            current += word + " "
                    if current.strip():
                        self.set_xy(LEFT_EDGE, y)
                        self.cell(TOTAL_WIDTH, 6, fa(current.strip()), align="R")
                        y += 6
                else:
                    self.set_xy(LEFT_EDGE, y)
                    self.cell(TOTAL_WIDTH, 6, fa(line), align="R")
                    y += 6
        else:
            if self.has_bold:
                self.set_font("Nazanin", "B", size=11)
            else:
                self.set_font("Nazanin", size=11)

            self.set_text_color(*GREEN)
            self.set_xy(LEFT_EDGE, y)
            self.cell(
                TOTAL_WIDTH, 7,
                fa("مغایرت گاری‌ها: بدون مغایرت"),
                align="R",
            )
            y += 7

        self.set_text_color(*BLACK)
        return y


# ==========================================
# توابع کمکی
# ==========================================

def _calculate_discrepancies(form, form_date):
    # داده‌های آلایشات از قبل گروه‌بندی شده‌اند
    grouped_alayeshat = form["alayeshat_rows"]
    
    total_assigned = 0.0
    for row in grouped_alayeshat:
        total_assigned += (
            safe_float(row.get("del_weight", 0)) +
            safe_float(row.get("jigar_weight", 0)) +
            safe_float(row.get("sangdane_weight", 0)) +
            safe_float(row.get("pa_weight", 0)) +
            safe_float(row.get("kale_weight", 0))
        )

    from functions import get_weighing_alayeshat_total_for_date
    total_weighing = get_weighing_alayeshat_total_for_date(form_date)
    alayeshat_diff = total_assigned - total_weighing

    piece_diffs = []
    for i, row in enumerate(form["morgh_rows"]):
        piece = safe_float(row.get("piece_weight", 0))
        sine = safe_float(row.get("sine_weight", 0))
        ran = safe_float(row.get("ran_weight", 0))
        bal = safe_float(row.get("bal_weight", 0))
        gardan = safe_float(row.get("gardan_weight", 0))

        total_parts = sine + ran + bal + gardan
        diff = piece - total_parts

        if piece == 0 and total_parts == 0:
            continue

        if abs(diff) > 0.01:
            piece_diffs.append(
                f"ردیف {i+1}: قطعه {piece:,.2f} ≠ جمع اجزا {total_parts:,.2f} (اختلاف {diff:+,.2f})"
            )

    return grouped_alayeshat, total_assigned, total_weighing, piece_diffs


def _calculate_cart_discrepancies(form, form_date):
    """
    محاسبه مغایرت گاری‌های فرم گزارش با گاری‌های فرم‌های توزین.

    سه حالت:
      الف) گاری انتخاب‌شده در گزارش که در فرم توزین نیست.
      ب) گاری موجود در فرم توزین که در گزارش تخصیص داده نشده.
      ج) گاری تکراری (در دو ردیف گزارش).

    خروجی: لیست رشته‌های مغایرت (خالی = بدون مغایرت)
    """
    from functions import get_available_carts_for_date

    # گاری‌های مجاز از فرم‌های توزین
    available_carts = get_available_carts_for_date(form_date)
    valid_cart_ids = {c["row_id"] for c in available_carts}
    num_by_id = {c["row_id"]: c["cart_number"] for c in available_carts}

    # گاری‌های انتخاب‌شده در گزارش
    selected_ids = []
    for row in form["packed_rows"]:
        for c in row.get("carts", []):
            cid = c.get("plucked_row_id")
            if cid:
                selected_ids.append(cid)

    cart_diffs = []

    # الف) گاری نامعتبر
    invalid = [cid for cid in selected_ids if cid not in valid_cart_ids]
    if invalid:
        invalid_str = "، ".join(str(cid) for cid in invalid)
        cart_diffs.append(
            f"گاری‌های انتخاب‌شده که در فرم توزین نیستند: {invalid_str}"
        )

    # ج) گاری تکراری
    seen = set()
    duplicates = []
    for cid in selected_ids:
        if cid in seen:
            duplicates.append(cid)
        seen.add(cid)

    if duplicates:
        dup_str = "، ".join(
            f"گاری {num_by_id.get(cid, cid)}" for cid in set(duplicates)
        )
        cart_diffs.append(f"گاری‌های تکراری در گزارش: {dup_str}")

    # ب) گاری توزین که تخصیص داده نشده
    unassigned = [cid for cid in valid_cart_ids if cid not in selected_ids]
    if unassigned:
        unassigned_str = "، ".join(
            f"گاری {num_by_id.get(cid, cid)}" for cid in unassigned
        )
        cart_diffs.append(
            f"گاری‌های فرم توزین که در گزارش تخصیص داده نشده‌اند: {unassigned_str}"
        )

    return cart_diffs


# ==========================================
# تولید PDF (بر اساس form_date)
# ==========================================

def generate_daily_pdf(form_date, output_path=None):
    from functions import get_daily_form_by_date

    form = get_daily_form_by_date(form_date)
    if not form:
        raise ValueError(f"فرم گزارش برای تاریخ {form_date} یافت نشد.")

    day_name = form.get("day_name", "")

    grouped_alayeshat, total_assigned, total_weighing, piece_diffs = _calculate_discrepancies(
        form, form_date
    )

    # محاسبه مغایرت گاری‌ها
    cart_diffs = _calculate_cart_discrepancies(form, form_date)

    pdf = DailyFormPDF()
    pdf.add_page()

    pdf.total_weighing_alayeshat = total_weighing
    pdf.piece_diffs = piece_diffs

    pdf.draw_header(form_date, day_name)

    y = pdf.get_y() + 6

    y = pdf.draw_packed_table(form["packed_rows"], y)

    y += 4
    y, alayeshat_totals = pdf.draw_alayeshat_table(grouped_alayeshat, y)

    y += 4
    y, morgh_totals = pdf.draw_morgh_table(form["morgh_rows"], y)

    y += 4
    y = pdf.draw_calc_and_discrepancies(
        alayeshat_totals=alayeshat_totals,
        morgh_totals=morgh_totals,
        waste=form["waste"],
        y=y,
    )

    # ── مغایرت گاری‌ها ──
    y = pdf.draw_cart_discrepancies(cart_diffs, y)

    if output_path is None:
        os.makedirs(REPORTS_DIR, exist_ok=True)
        safe_date = form_date.replace("/", "-")
        ts = datetime.now().strftime("%H%M%S")
        filename = f"daily_{safe_date}_{ts}.pdf"
        output_path = os.path.join(REPORTS_DIR, filename)

    pdf.output(output_path)
    return output_path


# ==========================================
# تولید Excel (xlsx) — بر اساس form_date
# ==========================================

def generate_daily_excel(form_date, output_path=None):
    from functions import get_daily_form_by_date

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ImportError(
            "کتابخانه openpyxl نصب نیست. لطفاً با دستور 'pip install openpyxl' نصب کنید."
        )

    form = get_daily_form_by_date(form_date)
    if not form:
        raise ValueError(f"فرم گزارش برای تاریخ {form_date} یافت نشد.")

    day_name = form.get("day_name", "")

    if output_path is None:
        os.makedirs(REPORTS_DIR, exist_ok=True)
        safe_date = form_date.replace("/", "-")
        ts = datetime.now().strftime("%H%M%S")
        filename = f"daily_{safe_date}_{ts}.xlsx"
        output_path = os.path.join(REPORTS_DIR, filename)

    grouped_alayeshat, total_assigned, total_weighing, piece_diffs = _calculate_discrepancies(
        form, form_date
    )

    # محاسبه مغایرت گاری‌ها
    cart_diffs = _calculate_cart_discrepancies(form, form_date)

    # ── استایل‌ها ──
    header_font = Font(name="B Nazanin", size=12, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1B9E8F", end_color="1B9E8F", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    title_font = Font(name="B Nazanin", size=16, bold=True, color="1B9E8F")
    subtitle_font = Font(name="B Nazanin", size=11, color="333333")
    section_font = Font(name="B Nazanin", size=13, bold=True, color="1B9E8F")

    cell_font = Font(name="B Nazanin", size=11, color="000000")
    cell_align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    total_font = Font(name="B Nazanin", size=11, bold=True, color="000000")
    total_fill = PatternFill(start_color="E0F2F1", end_color="E0F2F1", fill_type="solid")

    red_font = Font(name="B Nazanin", size=11, bold=True, color="C00000")
    green_font = Font(name="B Nazanin", size=11, bold=True, color="2E7D32")

    thin_border = Border(
        left=Side(style="thin", color="999999"),
        right=Side(style="thin", color="999999"),
        top=Side(style="thin", color="999999"),
        bottom=Side(style="thin", color="999999"),
    )

    # ── ساخت Workbook ──
    wb = Workbook()
    ws = wb.active
    ws.title = "گزارش روزانه"
    ws.sheet_view.rightToLeft = True

    # ── عنوان ──
    ws.merge_cells("A1:J1")
    c = ws["A1"]
    c.value = "فرم گزارش روزانه کشتارگاه"
    c.font = title_font
    c.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:J2")
    c = ws["A2"]
    date_text = f"تاریخ: {form_date}"
    if day_name:
        date_text += f"  ({day_name})"
    c.value = date_text
    c.font = subtitle_font
    c.alignment = Alignment(horizontal="center", vertical="center")

    current_row = 4

    # ══════════════════════════════════════
    # جدول ۱: مرغ پرکنده
    # ══════════════════════════════════════
    ws.merge_cells(f"A{current_row}:D{current_row}")
    c = ws[f"A{current_row}"]
    c.value = "مرغ پرکنده"
    c.font = section_font
    c.alignment = Alignment(horizontal="right", vertical="center")
    current_row += 1

    headers1 = ["ردیف", "نام خریدار", "شماره گاری", "وزن کل (kg)"]
    widths1 = [8, 30, 30, 20]

    for col_idx, (header, width) in enumerate(zip(headers1, widths1), start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    current_row += 1

    total_weight = 0.0
    for row in form["packed_rows"]:
        cart_numbers = ", ".join(c["cart_number"] for c in row.get("carts", []))
        weight = safe_float(row.get("total_net_weight", 0))
        total_weight += weight

        values = [row.get("row_num", ""), row.get("buyer_name", ""), cart_numbers, weight]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=value)
            cell.font = cell_font
            cell.alignment = cell_align_right if col_idx == 2 else cell_align_center
            cell.border = thin_border
        current_row += 1

    # جمع
    ws.cell(row=current_row, column=1, value="جمع")
    ws.cell(row=current_row, column=1).font = total_font
    ws.cell(row=current_row, column=1).fill = total_fill
    ws.cell(row=current_row, column=1).alignment = cell_align_center
    ws.cell(row=current_row, column=1).border = thin_border

    ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=3)
    c = ws.cell(row=current_row, column=2, value="")
    c.fill = total_fill
    c.border = thin_border

    c = ws.cell(row=current_row, column=4, value=total_weight)
    c.font = total_font
    c.fill = total_fill
    c.alignment = cell_align_center
    c.border = thin_border
    current_row += 2

    # ══════════════════════════════════════
    # جدول ۲: آلایشات
    # ══════════════════════════════════════
    ws.merge_cells(f"A{current_row}:G{current_row}")
    c = ws[f"A{current_row}"]
    c.value = "آلایشات"
    c.font = section_font
    c.alignment = Alignment(horizontal="right", vertical="center")
    current_row += 1

    headers2 = ["ردیف", "نام مشتری", "دل (kg)", "جگر (kg)", "سنگدان (kg)", "پا (kg)", "کله (kg)"]
    for col_idx, header in enumerate(headers2, start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
    current_row += 1

    total_del = total_jigar = total_sangdane = total_pa = total_kale = 0.0
    for row in grouped_alayeshat:
        del_w = safe_float(row.get("del_weight", 0))
        jigar_w = safe_float(row.get("jigar_weight", 0))
        sangdane_w = safe_float(row.get("sangdane_weight", 0))
        pa_w = safe_float(row.get("pa_weight", 0))
        kale_w = safe_float(row.get("kale_weight", 0))

        total_del += del_w
        total_jigar += jigar_w
        total_sangdane += sangdane_w
        total_pa += pa_w
        total_kale += kale_w

        values = [row.get("row_num", ""), row.get("customer_name", ""), del_w, jigar_w, sangdane_w, pa_w, kale_w]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=value)
            cell.font = cell_font
            cell.alignment = cell_align_right if col_idx == 2 else cell_align_center
            cell.border = thin_border
        current_row += 1

    values_total = ["جمع", "", total_del, total_jigar, total_sangdane, total_pa, total_kale]
    for col_idx, value in enumerate(values_total, start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=value)
        cell.font = total_font
        cell.fill = total_fill
        cell.alignment = cell_align_center
        cell.border = thin_border
    current_row += 2

    # ══════════════════════════════════════
    # جدول ۳: مرغ اصلاحی
    # ══════════════════════════════════════
    ws.merge_cells(f"A{current_row}:J{current_row}")
    c = ws[f"A{current_row}"]
    c.value = "مرغ اصلاحی"
    c.font = section_font
    c.alignment = Alignment(horizontal="right", vertical="center")
    current_row += 1

    headers3 = [
        "ردیف", "نام مشتری", "مرغ لکه (kg)", "قطعه (kg)", "سینه (kg)",
        "ران (kg)", "بال (kg)", "گردن (kg)", "مرغ کامل (kg)", "ضایعات (kg)"
    ]
    widths3 = [8, 25, 15, 15, 15, 15, 15, 15, 18, 15]

    for col_idx, (header, width) in enumerate(zip(headers3, widths3), start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    current_row += 1

    total_lake = total_piece = total_sine = total_ran = total_bal = total_gardan = total_whole = total_waste = 0.0

    for row in form["morgh_rows"]:
        lake_w = safe_float(row.get("lake_weight", 0))
        piece_w = safe_float(row.get("piece_weight", 0))
        sine_w = safe_float(row.get("sine_weight", 0))
        ran_w = safe_float(row.get("ran_weight", 0))
        bal_w = safe_float(row.get("bal_weight", 0))
        gardan_w = safe_float(row.get("gardan_weight", 0))
        whole_w = safe_float(row.get("whole_weight", 0))
        waste_w = safe_float(row.get("waste_weight", 0))

        total_lake += lake_w
        total_piece += piece_w
        total_sine += sine_w
        total_ran += ran_w
        total_bal += bal_w
        total_gardan += gardan_w
        total_whole += whole_w
        total_waste += waste_w

        values = [
            row.get("row_num", ""), row.get("customer_name", ""),
            lake_w, piece_w, sine_w, ran_w, bal_w, gardan_w, whole_w, waste_w,
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=value)
            cell.font = cell_font
            cell.alignment = cell_align_right if col_idx == 2 else cell_align_center
            cell.border = thin_border
        current_row += 1

    values_total = [
        "جمع", "",
        total_lake, total_piece, total_sine, total_ran, total_bal,
        total_gardan, total_whole, total_waste,
    ]
    for col_idx, value in enumerate(values_total, start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=value)
        cell.font = total_font
        cell.fill = total_fill
        cell.alignment = cell_align_center
        cell.border = thin_border
    current_row += 2

    # ══════════════════════════════════════
    # بخش محاسبه و مغایرت‌ها
    # ══════════════════════════════════════
    ws.merge_cells(f"A{current_row}:J{current_row}")
    c = ws[f"A{current_row}"]
    c.value = "محاسبه و مغایرت‌ها"
    c.font = section_font
    c.alignment = Alignment(horizontal="right", vertical="center")
    current_row += 1

    # محاسبه ضایعات
    calc_lines = [
        "(جمع مرغ کامل + جمع قطعه) - جمع مرغ لکه = ضایعات",
        f"({total_whole:,.2f} + {total_piece:,.2f}) - {total_lake:,.2f} = {total_waste:,.2f} کیلوگرم",
    ]
    for line in calc_lines:
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
        c = ws.cell(row=current_row, column=1, value=line)
        c.font = cell_font
        c.alignment = Alignment(horizontal="right", vertical="center")
        current_row += 1
    current_row += 1

    # مغایرت آلایشات
    diff_alayeshat = total_assigned - total_weighing
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
    c = ws.cell(row=current_row, column=1)
    if abs(diff_alayeshat) > 0.01:
        c.value = (
            f"مغایرت آلایشات: وزن جمعی آلایشات در فرم توزین برابر است با {total_weighing:,.2f} "
            f"و جمع وزنی آلایشات در اینجا برابر است با {total_assigned:,.2f} "
            f"(اختلاف: {diff_alayeshat:+,.2f} کیلوگرم)"
        )
        c.font = red_font
    else:
        c.value = f"مغایرت آلایشات: بدون مغایرت"
        c.font = green_font
    c.alignment = Alignment(horizontal="right", vertical="center")
    current_row += 1

    # مغایرت قطعه
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
    c = ws.cell(row=current_row, column=1)
    if piece_diffs:
        c.value = "مغایرت قطعه:\n" + "\n".join(piece_diffs)
        c.font = red_font
        lines_count = len(piece_diffs) + 1
        ws.row_dimensions[current_row].height = lines_count * 20
    else:
        c.value = "مغایرت قطعه: بدون مغایرت"
        c.font = green_font
    c.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    current_row += 1

    # ── مغایرت گاری‌ها (جدید) ──
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=10)
    c = ws.cell(row=current_row, column=1)
    if cart_diffs:
        c.value = "مغایرت گاری‌ها:\n" + "\n".join(cart_diffs)
        c.font = red_font
        lines_count = len(cart_diffs) + 1
        ws.row_dimensions[current_row].height = lines_count * 20
    else:
        c.value = "مغایرت گاری‌ها: بدون مغایرت"
        c.font = green_font
    c.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    current_row += 1

    # ── ذخیره ──
    wb.save(output_path)
    return output_path


# ==========================================
# تولید CSV (بر اساس form_date)
# ==========================================

def generate_daily_csv(form_date, output_path=None):
    from functions import get_daily_form_by_date

    form = get_daily_form_by_date(form_date)
    if not form:
        raise ValueError(f"فرم گزارش برای تاریخ {form_date} یافت نشد.")

    day_name = form.get("day_name", "")

    if output_path is None:
        os.makedirs(REPORTS_DIR, exist_ok=True)
        safe_date = form_date.replace("/", "-")
        ts = datetime.now().strftime("%H%M%S")
        filename = f"daily_{safe_date}_{ts}.csv"
        output_path = os.path.join(REPORTS_DIR, filename)

    grouped_alayeshat, total_assigned, total_weighing, piece_diffs = _calculate_discrepancies(
        form, form_date
    )

    # محاسبه مغایرت گاری‌ها
    cart_diffs = _calculate_cart_discrepancies(form, form_date)

    with open(output_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)

        writer.writerow(["فرم گزارش روزانه کشتارگاه"])
        writer.writerow([f"تاریخ: {form_date}", f"روز: {day_name}"])
        writer.writerow([])

        writer.writerow(["مرغ پرکنده"])
        writer.writerow(["ردیف", "نام خریدار", "شماره گاری", "وزن کل (kg)"])
        total_weight = 0.0
        for row in form["packed_rows"]:
            cart_numbers = ", ".join(c["cart_number"] for c in row.get("carts", []))
            weight = safe_float(row.get("total_net_weight", 0))
            total_weight += weight
            writer.writerow([row.get("row_num", ""), row.get("buyer_name", ""), cart_numbers, f"{weight:.2f}"])
        writer.writerow(["", "جمع", "", f"{total_weight:.2f}"])
        writer.writerow([])

        writer.writerow(["آلایشات"])
        writer.writerow(["ردیف", "نام مشتری", "دل (kg)", "جگر (kg)", "سنگدان (kg)", "پا (kg)", "کله (kg)"])
        total_del = total_jigar = total_sangdane = total_pa = total_kale = 0.0
        for row in grouped_alayeshat:
            del_w = safe_float(row.get("del_weight", 0))
            jigar_w = safe_float(row.get("jigar_weight", 0))
            sangdane_w = safe_float(row.get("sangdane_weight", 0))
            pa_w = safe_float(row.get("pa_weight", 0))
            kale_w = safe_float(row.get("kale_weight", 0))

            total_del += del_w
            total_jigar += jigar_w
            total_sangdane += sangdane_w
            total_pa += pa_w
            total_kale += kale_w

            writer.writerow([
                row.get("row_num", ""), row.get("customer_name", ""),
                f"{del_w:.2f}" if del_w else "",
                f"{jigar_w:.2f}" if jigar_w else "",
                f"{sangdane_w:.2f}" if sangdane_w else "",
                f"{pa_w:.2f}" if pa_w else "",
                f"{kale_w:.2f}" if kale_w else "",
            ])
        writer.writerow(["", "جمع", f"{total_del:.2f}", f"{total_jigar:.2f}", f"{total_sangdane:.2f}", f"{total_pa:.2f}", f"{total_kale:.2f}"])
        writer.writerow([])

        writer.writerow(["مرغ اصلاحی"])
        writer.writerow([
            "ردیف", "نام مشتری", "مرغ لکه (kg)", "قطعه (kg)", "سینه (kg)",
            "ران (kg)", "بال (kg)", "گردن (kg)", "مرغ کامل (kg)", "ضایعات (kg)"
        ])

        total_lake = total_piece = total_sine = total_ran = total_bal = total_gardan = total_whole = total_waste = 0.0

        for row in form["morgh_rows"]:
            lake_w = safe_float(row.get("lake_weight", 0))
            piece_w = safe_float(row.get("piece_weight", 0))
            sine_w = safe_float(row.get("sine_weight", 0))
            ran_w = safe_float(row.get("ran_weight", 0))
            bal_w = safe_float(row.get("bal_weight", 0))
            gardan_w = safe_float(row.get("gardan_weight", 0))
            whole_w = safe_float(row.get("whole_weight", 0))
            waste_w = safe_float(row.get("waste_weight", 0))

            total_lake += lake_w
            total_piece += piece_w
            total_sine += sine_w
            total_ran += ran_w
            total_bal += bal_w
            total_gardan += gardan_w
            total_whole += whole_w
            total_waste += waste_w

            writer.writerow([
                row.get("row_num", ""), row.get("customer_name", ""),
                f"{lake_w:.2f}" if lake_w else "",
                f"{piece_w:.2f}" if piece_w else "",
                f"{sine_w:.2f}" if sine_w else "",
                f"{ran_w:.2f}" if ran_w else "",
                f"{bal_w:.2f}" if bal_w else "",
                f"{gardan_w:.2f}" if gardan_w else "",
                f"{whole_w:.2f}" if whole_w else "",
                f"{waste_w:.2f}" if waste_w else "",
            ])

        writer.writerow([
            "", "جمع",
            f"{total_lake:.2f}", f"{total_piece:.2f}", f"{total_sine:.2f}",
            f"{total_ran:.2f}", f"{total_bal:.2f}", f"{total_gardan:.2f}",
            f"{total_whole:.2f}", f"{total_waste:.2f}"
        ])
        writer.writerow([])

        writer.writerow(["محاسبه"])
        writer.writerow(["(جمع مرغ کامل + جمع قطعه) - جمع مرغ لکه = ضایعات"])
        writer.writerow([f"({total_whole:.2f} + {total_piece:.2f}) - {total_lake:.2f} = {total_waste:.2f} کیلوگرم"])
        writer.writerow([])

        writer.writerow(["مغایرت‌ها"])

        diff_alayeshat = total_assigned - total_weighing
        if abs(diff_alayeshat) > 0.01:
            writer.writerow([
                f"مغایرت آلایشات: جمع تخصیص {total_assigned:.2f} - جمع فرم توزین {total_weighing:.2f} = {diff_alayeshat:+.2f} کیلوگرم"
            ])
        else:
            writer.writerow([f"مغایرت آلایشات: بدون مغایرت (جمع {total_assigned:.2f} کیلوگرم)"])

        if piece_diffs:
            for line in piece_diffs:
                writer.writerow([line])
        else:
            writer.writerow(["مغایرت قطعه: بدون مغایرت"])

        # ── مغایرت گاری‌ها (جدید) ──
        if cart_diffs:
            for line in cart_diffs:
                writer.writerow([line])
        else:
            writer.writerow(["مغایرت گاری‌ها: بدون مغایرت"])

    return output_path


# ==========================================
# تست
# ==========================================

if __name__ == "__main__":
    from database import initialize_database

    initialize_database()

    form_date = "1405/07/30"

    try:
        path = generate_daily_pdf(form_date)
        print(f"✅ PDF ساخته شد: {path}")

        try:
            os.startfile(path)
        except Exception:
            pass
    except Exception as e:
        print(f"❌ خطا: {e}")
