# -*- coding: utf-8 -*-
"""
تولید PDF و Excel برای گزارش‌گیری جامع (چند بخشی)
"""
import os
from datetime import datetime
from fpdf import FPDF
import arabic_reshaper
from bidi.algorithm import get_display

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
FONT_REG_PATH = os.path.join(BASE_DIR, "assets", "fonts", "BNazanin.ttf")
REPORTS_DIR = os.path.join(BASE_DIR, "reports")

def fa(text):
    if not text: return ""
    reshaped = arabic_reshaper.reshape(str(text))
    return get_display(reshaped)

def fa_digits(text):
    if text is None: return ""
    s = str(text)
    mapping = {'0': '۰', '1': '۱', '2': '۲', '3': '۳', '4': '۴', '5': '۵', '6': '۶', '7': '۷', '8': '۸', '9': '۹', '.': '٫', ',': '٬'}
    return "".join(mapping.get(c, c) for c in s)

def mixed_fa(text):
    if not text: return ""
    import re
    parts = re.split(r'([۰-۹0-9,.٬٫]+)', str(text))
    result = []
    for part in parts:
        if re.match(r'^[۰-۹0-9,.٬٫]+$', part):
            result.append(fa_digits(part))
        else:
            if part.strip():
                result.append(fa(part))
    return "".join(result)

class AdminReportPDF(FPDF):
    def __init__(self):
        super().__init__(orientation="P", unit="mm", format="A4")
        self.set_auto_page_break(auto=True, margin=15)
        self.set_margins(15, 15, 15)
        if not os.path.exists(FONT_REG_PATH):
            raise FileNotFoundError(f"فونت پیدا نشد: {FONT_REG_PATH}")
        self.add_font("Nazanin", "", FONT_REG_PATH)

    def header(self):
        self.set_font("Nazanin", size=16)
        self.set_text_color(20, 122, 111)
        self.cell(0, 10, fa("سیستم توزین محصولات"), align="C", new_x="LMARGIN", new_y="NEXT")
        self.set_font("Nazanin", size=12)
        self.set_text_color(100, 100, 100)
        self.cell(0, 8, fa("گزارش‌گیری جامع"), align="C", new_x="LMARGIN", new_y="NEXT")
        self.ln(5)
        self.set_draw_color(20, 122, 111)
        self.line(15, self.get_y(), 195, self.get_y())
        self.ln(5)

    def footer(self):
        self.set_y(-15)
        self.set_font("Nazanin", size=8)
        self.set_text_color(150, 150, 150)
        self.cell(0, 10, fa(f"صفحه {self.page_no()}"), align="C")

    def section_title(self, title):
        self.set_font("Nazanin", size=14)
        self.set_text_color(20, 122, 111)
        self.ln(5)
        self.cell(0, 10, fa(title), align="R", new_x="LMARGIN", new_y="NEXT")
        self.ln(2)


def generate_admin_report_pdf(summary, details, daily, operators, customers, products, date_from, date_to, output_path=None):
    os.makedirs(REPORTS_DIR, exist_ok=True)
    if output_path is None:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_from = date_from.replace("/", "-")
        safe_to = date_to.replace("/", "-")
        filename = f"admin_report_{safe_from}_{safe_to}_{ts}.pdf"
        output_path = os.path.join(REPORTS_DIR, filename)

    pdf = AdminReportPDF()
    pdf.add_page()

    pdf.set_font("Nazanin", size=11)
    pdf.set_text_color(50, 50, 50)
    pdf.cell(0, 8, mixed_fa(f"بازه گزارش: {date_from} تا {date_to}"), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    # ۱. خلاصه
    pdf.section_title("خلاصه آمار")
    pdf.set_font("Nazanin", size=11)
    pdf.set_text_color(40, 40, 40)
    
    summary_items = [
        ("تعداد کشتارها", fa_digits(summary["slaughter_count"])),
        ("مجموع پای قلاب", fa_digits(summary["total_hook"])),
        ("مجموع وزن خالص توزین", mixed_fa(f"{summary['total_net_weight']:,.0f}") + " " + fa("کیلوگرم")),
        ("تعداد قبض‌ها", fa_digits(summary["receipt_count"])),
        ("مجموع وزن قبض‌ها", mixed_fa(f"{summary['receipt_weight']:,.0f}") + " " + fa("کیلوگرم")),
        ("تعداد صورتحساب‌ها", fa_digits(summary["invoice_count"])),
        ("جمع کل مبالغ صورتحساب", mixed_fa(f"{summary['invoice_total']:,.0f}") + " " + fa("ریال")),
    ]
    
    for label, value in summary_items:
        pdf.cell(60, 9, fa(label), border=1, align="R")
        pdf.cell(120, 9, str(value), border=1, align="R", new_x="LMARGIN", new_y="NEXT")

    # ۲. تفکیکی کشتارها
    # ۲. تفکیکی کشتارها
    pdf.section_title("گزارش تفکیکی کشتارها")
    if not details:
        pdf.cell(0, 10, fa("هیچ کشتاری در این بازه یافت نشد."), align="C")
    else:
        pdf.set_font("Nazanin", size=9)
        pdf.set_fill_color(20, 122, 111)
        pdf.set_text_color(255, 255, 255)
        
        headers = ["تاریخ", "کشتار", "پای قلاب", "وزن خالص", "قبض‌ها", "صورتحساب", "گاری", "پرکنده", "آلایشات"]
        # مجموع عرض‌ها دقیقاً 180 میلی‌متر (برابر با عرض سایر جداول)
        widths = [22, 16, 16, 20, 16, 24, 14, 26, 26]
        
        for h, w in zip(headers, widths):
            pdf.cell(w, 8, fa(h), border=1, align="C", fill=True)
        pdf.ln()

        pdf.set_font("Nazanin", size=8)
        pdf.set_text_color(40, 40, 40)
        
        for d in details:
            pdf.cell(widths[0], 7, fa_digits(d["form_date"]), border=1, align="C")
            pdf.cell(widths[1], 7, fa_digits(d["slaughter_number"]), border=1, align="C")
            pdf.cell(widths[2], 7, fa_digits(d["hook_count"]), border=1, align="C")
            pdf.cell(widths[3], 7, fa_digits(f"{d['net_weight']:,.0f}"), border=1, align="C")
            pdf.cell(widths[4], 7, fa_digits(d["receipt_count"]), border=1, align="C")
            pdf.cell(widths[5], 7, fa_digits(f"{d['invoice_total']:,.0f}"), border=1, align="C")
            pdf.cell(widths[6], 7, fa_digits(d["cart_count"]), border=1, align="C")
            pdf.cell(widths[7], 7, fa_digits(f"{d['plucked_weight']:,.0f}"), border=1, align="C")
            pdf.cell(widths[8], 7, fa_digits(f"{d['alayeshat_weight']:,.0f}"), border=1, align="C")
            pdf.ln()
            
    # ۳. تفکیکی روزانه
    pdf.section_title("گزارش تفکیکی روزانه")
    if not daily:
        pdf.cell(0, 10, fa("هیچ داده‌ای نیست."), align="C")
    else:
        pdf.set_font("Nazanin", size=10)
        pdf.set_fill_color(20, 122, 111)
        pdf.set_text_color(255, 255, 255)
        
        headers = ["تاریخ", "روز", "تعداد کشتار", "پای قلاب", "وزن خالص"]
        widths = [35, 35, 40, 35, 35]
        for h, w in zip(headers, widths):
            pdf.cell(w, 8, fa(h), border=1, align="C", fill=True)
        pdf.ln()

        pdf.set_font("Nazanin", size=9)
        pdf.set_text_color(40, 40, 40)
        for r in daily:
            pdf.cell(widths[0], 7, fa_digits(r["form_date"]), border=1, align="C")
            pdf.cell(widths[1], 7, fa(r["day_name"]), border=1, align="C")
            pdf.cell(widths[2], 7, fa_digits(r["slaughter_count"]), border=1, align="C")
            pdf.cell(widths[3], 7, fa_digits(r["total_hook"]), border=1, align="C")
            pdf.cell(widths[4], 7, fa_digits(f"{r['total_net']:,.0f}"), border=1, align="C")
            pdf.ln()

    # ۴. تفکیکی اپراتورها
    pdf.section_title("گزارش تفکیکی اپراتورها")
    if not operators:
        pdf.cell(0, 10, fa("هیچ داده‌ای نیست."), align="C")
    else:
        pdf.set_font("Nazanin", size=10)
        pdf.set_fill_color(20, 122, 111)
        pdf.set_text_color(255, 255, 255)
        
        headers = ["نام کاربر", "تعداد کشتار", "پای قلاب", "وزن خالص"]
        widths = [60, 40, 40, 40]
        for h, w in zip(headers, widths):
            pdf.cell(w, 8, fa(h), border=1, align="C", fill=True)
        pdf.ln()

        pdf.set_font("Nazanin", size=9)
        pdf.set_text_color(40, 40, 40)
        for r in operators:
            pdf.cell(widths[0], 7, fa(r["full_name"]), border=1, align="R")
            pdf.cell(widths[1], 7, fa_digits(r["slaughter_count"]), border=1, align="C")
            pdf.cell(widths[2], 7, fa_digits(r["total_hook"]), border=1, align="C")
            pdf.cell(widths[3], 7, fa_digits(f"{r['total_net']:,.0f}"), border=1, align="C")
            pdf.ln()

    # ۵. تفکیکی مشتری‌ها
    pdf.section_title("گزارش تفکیکی مشتری‌ها")
    if not customers["buyers"] and not customers["alayeshat"]:
        pdf.cell(0, 10, fa("هیچ داده‌ای نیست."), align="C")
    else:
        if customers["buyers"]:
            pdf.set_font("Nazanin", size=11)
            pdf.set_text_color(20, 122, 111)
            pdf.cell(0, 8, fa("خریداران (مرغ پرکنده):"), align="R", new_x="LMARGIN", new_y="NEXT")
            
            pdf.set_font("Nazanin", size=9)
            pdf.set_fill_color(20, 122, 111)
            pdf.set_text_color(255, 255, 255)
            headers = ["نام خریدار", "تعداد ردیف", "وزن کل"]
            widths = [90, 45, 45]
            for h, w in zip(headers, widths):
                pdf.cell(w, 8, fa(h), border=1, align="C", fill=True)
            pdf.ln()

            pdf.set_font("Nazanin", size=9)
            pdf.set_text_color(40, 40, 40)
            for r in customers["buyers"]:
                pdf.cell(widths[0], 7, fa(r["name"]), border=1, align="R")
                pdf.cell(widths[1], 7, fa_digits(r["row_count"]), border=1, align="C")
                pdf.cell(widths[2], 7, fa_digits(f"{r['total_weight']:,.0f}"), border=1, align="C")
                pdf.ln()

        if customers["alayeshat"]:
            pdf.ln(3)
            pdf.set_font("Nazanin", size=11)
            pdf.set_text_color(20, 122, 111)
            pdf.cell(0, 8, fa("مشتریان آلایشات:"), align="R", new_x="LMARGIN", new_y="NEXT")
            
            pdf.set_font("Nazanin", size=9)
            pdf.set_fill_color(20, 122, 111)
            pdf.set_text_color(255, 255, 255)
            headers = ["نام مشتری", "تعداد ردیف", "وزن کل"]
            widths = [90, 45, 45]
            for h, w in zip(headers, widths):
                pdf.cell(w, 8, fa(h), border=1, align="C", fill=True)
            pdf.ln()

            pdf.set_font("Nazanin", size=9)
            pdf.set_text_color(40, 40, 40)
            for r in customers["alayeshat"]:
                pdf.cell(widths[0], 7, fa(r["name"]), border=1, align="R")
                pdf.cell(widths[1], 7, fa_digits(r["row_count"]), border=1, align="C")
                pdf.cell(widths[2], 7, fa_digits(f"{r['total_weight']:,.0f}"), border=1, align="C")
                pdf.ln()

    # ۶. تفکیکی محصولات
    pdf.section_title("گزارش تفکیکی محصولات (آلایشات)")
    if not products:
        pdf.cell(0, 10, fa("هیچ داده‌ای نیست."), align="C")
    else:
        pdf.set_font("Nazanin", size=10)
        pdf.set_fill_color(20, 122, 111)
        pdf.set_text_color(255, 255, 255)
        
        headers = ["نام محصول", "تعداد ردیف", "وزن کل"]
        widths = [90, 45, 45]
        for h, w in zip(headers, widths):
            pdf.cell(w, 8, fa(h), border=1, align="C", fill=True)
        pdf.ln()

        pdf.set_font("Nazanin", size=9)
        pdf.set_text_color(40, 40, 40)
        for r in products:
            pdf.cell(widths[0], 7, fa(r["product_name"]), border=1, align="R")
            pdf.cell(widths[1], 7, fa_digits(r["row_count"]), border=1, align="C")
            pdf.cell(widths[2], 7, fa_digits(f"{r['total_weight']:,.0f}"), border=1, align="C")
            pdf.ln()

    pdf.output(output_path)
    return output_path


def generate_admin_report_excel(summary, details, daily, operators, customers, products, date_from, date_to, output_path=None):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ImportError("کتابخانه openpyxl نصب نیست. لطفاً با دستور 'pip install openpyxl' نصب کنید.")

    os.makedirs(REPORTS_DIR, exist_ok=True)
    if output_path is None:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_from = date_from.replace("/", "-")
        safe_to = date_to.replace("/", "-")
        filename = f"admin_report_{safe_from}_{safe_to}_{ts}.xlsx"
        output_path = os.path.join(REPORTS_DIR, filename)

    wb = Workbook()
    
    title_font = Font(name="B Nazanin", size=16, bold=True, color="1B9E8F")
    section_font = Font(name="B Nazanin", size=14, bold=True, color="1B9E8F")
    header_font = Font(name="B Nazanin", size=12, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1B9E8F", end_color="1B9E8F", fill_type="solid")
    cell_font = Font(name="B Nazanin", size=11, color="000000")
    total_font = Font(name="B Nazanin", size=11, bold=True, color="000000")
    
    border = Border(
        left=Side(style="thin", color="999999"), right=Side(style="thin", color="999999"),
        top=Side(style="thin", color="999999"), bottom=Side(style="thin", color="999999"),
    )

    # ── شیت ۱: خلاصه ──
    ws = wb.active
    ws.title = "خلاصه"
    ws.sheet_view.rightToLeft = True

    ws.merge_cells("A1:D1")
    ws["A1"].value = "گزارش‌گیری جامع — خلاصه"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:D2")
    ws["A2"].value = f"بازه گزارش: {date_from} تا {date_to}"
    ws["A2"].font = Font(name="B Nazanin", size=11, color="333333")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    row = 4
    summary_items = [
        ("تعداد کشتارها", summary["slaughter_count"]),
        ("مجموع پای قلاب", summary["total_hook"]),
        ("مجموع وزن خالص توزین", f"{summary['total_net_weight']:,.0f} کیلوگرم"),
        ("تعداد قبض‌ها", summary["receipt_count"]),
        ("مجموع وزن قبض‌ها", f"{summary['receipt_weight']:,.0f} کیلوگرم"),
        ("تعداد صورتحساب‌ها", summary["invoice_count"]),
        ("جمع کل مبالغ صورتحساب", f"{summary['invoice_total']:,.0f} ریال"),
    ]
    
    for label, value in summary_items:
        ws.cell(row=row, column=1, value=label).font = cell_font
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="right")
        ws.cell(row=row, column=1).border = border
        
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
        ws.cell(row=row, column=2, value=value).font = total_font
        ws.cell(row=row, column=2).alignment = Alignment(horizontal="right")
        ws.cell(row=row, column=2).border = border
        row += 1

    # ── شیت ۲: کشتارها ──
    ws2 = wb.create_sheet("کشتارها")
    ws2.sheet_view.rightToLeft = True

    headers2 = ["تاریخ", "کشتار", "پای قلاب", "وزن خالص", "قبض‌ها", "صورتحساب", "گاری", "پرکنده", "آلایشات", "اپراتور"]
    widths2 = [15, 12, 12, 15, 12, 18, 10, 15, 15, 20]
    
    for col_idx, (h, w) in enumerate(zip(headers2, widths2), start=1):
        cell = ws2.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        ws2.column_dimensions[get_column_letter(col_idx)].width = w

    for row_idx, d in enumerate(details, start=2):
        values = [
            d["form_date"], d["slaughter_number"], d["hook_count"], f"{d['net_weight']:,.0f}",
            d["receipt_count"], f"{d['invoice_total']:,.0f}", d["cart_count"],
            f"{d['plucked_weight']:,.0f}", f"{d['alayeshat_weight']:,.0f}", d["operator_name"],
        ]
        for col_idx, val in enumerate(values, start=1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            cell.font = cell_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

    # ── شیت ۳: روزانه ──
    ws3 = wb.create_sheet("روزانه")
    ws3.sheet_view.rightToLeft = True

    headers3 = ["تاریخ", "روز", "تعداد کشتار", "پای قلاب", "وزن خالص"]
    for col_idx, h in enumerate(headers3, start=1):
        cell = ws3.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        ws3.column_dimensions[get_column_letter(col_idx)].width = 18

    for row_idx, r in enumerate(daily, start=2):
        values = [r["form_date"], r["day_name"], r["slaughter_count"], r["total_hook"], f"{r['total_net']:,.0f}"]
        for col_idx, val in enumerate(values, start=1):
            cell = ws3.cell(row=row_idx, column=col_idx, value=val)
            cell.font = cell_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

    # ── شیت ۴: اپراتورها ──
    ws4 = wb.create_sheet("اپراتورها")
    ws4.sheet_view.rightToLeft = True

    headers4 = ["نام کاربر", "نام کاربری", "تعداد کشتار", "پای قلاب", "وزن خالص"]
    for col_idx, h in enumerate(headers4, start=1):
        cell = ws4.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        ws4.column_dimensions[get_column_letter(col_idx)].width = 20

    for row_idx, r in enumerate(operators, start=2):
        values = [r["full_name"], r["username"], r["slaughter_count"], r["total_hook"], f"{r['total_net']:,.0f}"]
        for col_idx, val in enumerate(values, start=1):
            cell = ws4.cell(row=row_idx, column=col_idx, value=val)
            cell.font = cell_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

    # ── شیت ۵: مشتری‌ها ──
    ws5 = wb.create_sheet("مشتری‌ها")
    ws5.sheet_view.rightToLeft = True

    row = 1
    ws5.cell(row=row, column=1, value="خریداران (مرغ پرکنده)").font = section_font
    row += 1
    headers5a = ["نام خریدار", "تعداد ردیف", "وزن کل"]
    for col_idx, h in enumerate(headers5a, start=1):
        cell = ws5.cell(row=row, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        ws5.column_dimensions[get_column_letter(col_idx)].width = 25
    row += 1
    for r in customers["buyers"]:
        values = [r["name"], r["row_count"], f"{r['total_weight']:,.0f}"]
        for col_idx, val in enumerate(values, start=1):
            cell = ws5.cell(row=row, column=col_idx, value=val)
            cell.font = cell_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        row += 1

    row += 2
    ws5.cell(row=row, column=1, value="مشتریان آلایشات").font = section_font
    row += 1
    headers5b = ["نام مشتری", "تعداد ردیف", "وزن کل"]
    for col_idx, h in enumerate(headers5b, start=1):
        cell = ws5.cell(row=row, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    row += 1
    for r in customers["alayeshat"]:
        values = [r["name"], r["row_count"], f"{r['total_weight']:,.0f}"]
        for col_idx, val in enumerate(values, start=1):
            cell = ws5.cell(row=row, column=col_idx, value=val)
            cell.font = cell_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        row += 1

    # ── شیت ۶: محصولات ──
    ws6 = wb.create_sheet("محصولات")
    ws6.sheet_view.rightToLeft = True

    headers6 = ["نام محصول", "تعداد ردیف", "وزن کل"]
    for col_idx, h in enumerate(headers6, start=1):
        cell = ws6.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        ws6.column_dimensions[get_column_letter(col_idx)].width = 25

    for row_idx, r in enumerate(products, start=2):
        values = [r["product_name"], r["row_count"], f"{r['total_weight']:,.0f}"]
        for col_idx, val in enumerate(values, start=1):
            cell = ws6.cell(row=row_idx, column=col_idx, value=val)
            cell.font = cell_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

    wb.save(output_path)
    return output_path
