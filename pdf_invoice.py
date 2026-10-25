# -*- coding: utf-8 -*-
"""
تولید PDF و Excel برای صورتحساب کشتار روزانه
"""
import os
import textwrap
from fpdf import FPDF
import arabic_reshaper
from bidi.algorithm import get_display

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
FONT_REG_PATH = os.path.join(BASE_DIR, "assets", "fonts", "BNazanin.ttf")
FONT_BOLD_PATH = os.path.join(BASE_DIR, "assets", "fonts", "BNazaninBold.ttf")
REPORTS_DIR = os.path.join(BASE_DIR, "reports")


def fa(text):
    if not text:
        return ""
    reshaped = arabic_reshaper.reshape(str(text))
    return get_display(reshaped)


def _num(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


# ==========================================
# PDF
# ==========================================

def generate_daily_invoice_pdf(form_date, output_path=None):
    from functions import get_daily_invoice

    invoice = get_daily_invoice(form_date)
    if not invoice:
        raise ValueError(f"صورتحساب روز {form_date} یافت نشد.")

    os.makedirs(REPORTS_DIR, exist_ok=True)
    if output_path is None:
        safe_date = form_date.replace("/", "-")
        output_path = os.path.join(REPORTS_DIR, f"invoice_{safe_date}.pdf")

    class PDF(FPDF):
        def __init__(self):
            super().__init__(orientation="P", unit="mm", format="A4")
            self.set_auto_page_break(auto=False)
            self.set_margins(10, 10, 10)
            self.add_font("Nazanin", "", FONT_REG_PATH)
            if os.path.exists(FONT_BOLD_PATH):
                self.add_font("Nazanin", "B", FONT_BOLD_PATH)
                self.has_bold = True
            else:
                self.has_bold = False

        def header(self):
            if self.has_bold:
                self.set_font("Nazanin", "B", size=16)
            else:
                self.set_font("Nazanin", size=16)
            self.set_text_color(20, 122, 111)
            self.cell(0, 10, fa("صورتحساب کشتار روزانه"), align="C", new_x="LMARGIN", new_y="NEXT")
            self.set_font("Nazanin", size=11)
            self.set_text_color(80, 80, 80)
            self.cell(0, 6, fa(f"تاریخ: {form_date}"), align="C", new_x="LMARGIN", new_y="NEXT")
            self.ln(3)
            self.set_draw_color(20, 122, 111)
            self.set_line_width(0.5)
            self.line(10, self.get_y(), 200, self.get_y())
            self.ln(4)

        def footer(self):
            self.set_y(-10)
            self.set_font("Nazanin", size=8)
            self.set_text_color(150, 150, 150)
            self.cell(0, 6, fa("سیستم توزین محصولات"), align="C")

    pdf = PDF()
    pdf.add_page()

    pdf.set_font("Nazanin", size=10)
    pdf.set_text_color(40, 40, 40)
    inv_num = invoice.get("invoice_number") or "—"
    inv_date = invoice.get("invoice_date") or "—"

    pdf.cell(90, 7, fa(f"شماره صورتحساب: {inv_num}"), align="R")
    pdf.cell(90, 7, fa(f"تاریخ: {inv_date}"), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    # ── جدول وزن‌کشی ──
    if pdf.has_bold:
        pdf.set_font("Nazanin", "B", size=11)
    else:
        pdf.set_font("Nazanin", size=11)
    pdf.set_text_color(20, 122, 111)
    pdf.cell(0, 7, fa("جدول وزن‌کشی (کیلوگرم)"), align="R", new_x="LMARGIN", new_y="NEXT")

    widths = [10, 22, 22, 22, 20, 25, 30]
    headers = ["ردیف", "وزن پر", "وزن خالی", "خالص", "تعداد", "فی", "مبلغ"]

    if pdf.has_bold:
        pdf.set_font("Nazanin", "B", size=9)
    else:
        pdf.set_font("Nazanin", size=9)
    pdf.set_text_color(255, 255, 255)
    pdf.set_fill_color(20, 122, 111)

    for h, w in zip(headers, widths):
        pdf.cell(w, 7, fa(h), border=1, align="C", fill=True)
    pdf.ln()

    pdf.set_font("Nazanin", size=9)
    pdf.set_text_color(40, 40, 40)
    total_weight = 0
    total_amount = 0
    for i in range(1, 6):
        full_val = _num(invoice.get(f"w{i}_full"))
        empty_val = _num(invoice.get(f"w{i}_empty"))
        net_val = max(0, full_val - empty_val)
        qty_val = _num(invoice.get(f"w{i}_qty"))
        rate_val = _num(invoice.get(f"w{i}_rate"))
        amt_val = net_val * (qty_val if qty_val > 0 else 1) * rate_val

        vals = [
            str(i),
            f"{full_val:,.2f}",
            f"{empty_val:,.2f}",
            f"{net_val:,.2f}",
            f"{qty_val:,.0f}",
            f"{rate_val:,.0f}",
            f"{amt_val:,.0f}",
        ]
        for val, w in zip(vals, widths):
            pdf.cell(w, 7, val, border=1, align="C")
        pdf.ln()
        total_weight += net_val
        total_amount += amt_val

    # جمع وزن‌کشی (بدون فاصله)
    if pdf.has_bold:
        pdf.set_font("Nazanin", "B", size=9)
    else:
        pdf.set_font("Nazanin", size=9)
    pdf.set_fill_color(235, 245, 245)
    pdf.set_text_color(20, 122, 111)
    pdf.cell(widths[0] + widths[1] + widths[2], 7, fa("جمع"), border=1, align="C", fill=True)
    pdf.cell(widths[3], 7, f"{total_weight:,.2f}", border=1, align="C", fill=True)
    pdf.cell(widths[4] + widths[5], 7, "", border=1, align="C", fill=True)
    pdf.cell(widths[6], 7, f"{total_amount:,.0f}", border=1, align="C", fill=True)
    pdf.ln()

    # ── کشتار و بسته‌بندی (چسبیده به جدول بالا) ──
    if pdf.has_bold:
        pdf.set_font("Nazanin", "B", size=9)
    else:
        pdf.set_font("Nazanin", size=9)
    pdf.set_fill_color(235, 245, 245)
    pdf.set_text_color(20, 122, 111)
    pdf.cell(sum(widths), 7, fa("کشتار و بسته‌بندی"), border=1, align="C", fill=True)
    pdf.ln()

    pdf.set_font("Nazanin", size=9)
    pdf.set_text_color(40, 40, 40)

    c_label = widths[0] + widths[1]
    c_amt   = widths[2]
    c_rate  = widths[3]
    c_total = widths[4] + widths[5] + widths[6]

    rows = [
        ("کشتار مرغ گوشتی", _num(invoice.get("c1_amount")), _num(invoice.get("c1_rate")), _num(invoice.get("c1_total"))),
        ("بسته بندی مرغ گوشتی", _num(invoice.get("c2_amount")), _num(invoice.get("c2_rate")), _num(invoice.get("c2_total"))),
    ]
    for label, amt, rate, total in rows:
        pdf.cell(c_label, 7, fa(label), border=1, align="R")
        pdf.cell(c_amt, 7, f"{amt:,.2f}", border=1, align="C")
        pdf.cell(c_rate, 7, f"{rate:,.0f}", border=1, align="C")
        pdf.cell(c_total, 7, f"{total:,.0f}", border=1, align="C")
        pdf.ln()

    pdf.ln(6)

    # ── خلاصه فاکتور اول ──
    if pdf.has_bold:
        pdf.set_font("Nazanin", "B", size=11)
    else:
        pdf.set_font("Nazanin", size=11)
    pdf.set_text_color(20, 122, 111)
    pdf.cell(0, 7, fa("خلاصه فاکتور (کشتار و بسته‌بندی)"), align="R", new_x="LMARGIN", new_y="NEXT")

    pdf.set_font("Nazanin", size=10)
    pdf.set_text_color(40, 40, 40)
    base = _num(invoice.get("c1_total")) + _num(invoice.get("c2_total"))
    vat = _num(invoice.get("vat_amount"))
    grand = _num(invoice.get("grand_total"))

    pdf.cell(0, 7, fa(f"مبلغ پایه: {base:,.0f} ریال"), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 7, fa(f"ارزش افزوده (۱۰٪): {vat:,.0f} ریال"), align="R", new_x="LMARGIN", new_y="NEXT")
    if pdf.has_bold:
        pdf.set_font("Nazanin", "B", size=11)
    else:
        pdf.set_font("Nazanin", size=11)
    pdf.cell(0, 7, fa(f"جمع کل فاکتور: {grand:,.0f} ریال"), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)

    # ── هزینه‌ها ──
    if pdf.has_bold:
        pdf.set_font("Nazanin", "B", size=11)
    else:
        pdf.set_font("Nazanin", size=11)
    pdf.set_text_color(20, 122, 111)
    pdf.cell(0, 7, fa("هزینه‌ها"), align="R", new_x="LMARGIN", new_y="NEXT")

    expense_names = [
        "هزینه یخ مازاد داخل خودرو",
        "بسته بندی دل و کگر و سنگدان",
        "قطعه بندی مرغ",
        "پاکسازی دل و کگر و سنگدان",
        "پوست گیری پای مرغ",
        "صدور مجوز",
        "پالمپ",
        "عوارض دامپزشکی"
    ]

    # ── سربرگ جدول هزینه‌ها ──
    if pdf.has_bold:
        pdf.set_font("Nazanin", "B", size=9)
    else:
        pdf.set_font("Nazanin", size=9)
    pdf.set_text_color(255, 255, 255)
    pdf.set_fill_color(20, 122, 111)
    exp_widths = [60, 30, 30, 40]
    exp_headers = ["شرح هزینه", "وزن (کیلوگرم)", "فی", "مبلغ"]
    for h, w in zip(exp_headers, exp_widths):
        pdf.cell(w, 7, fa(h), border=1, align="C", fill=True)
    pdf.ln()

    # ── ردیف‌های هزینه ──
    pdf.set_font("Nazanin", size=9)
    pdf.set_text_color(40, 40, 40)
    e_taxable_total = 0.0
    e_avarez_total = 0.0
    for i, name in enumerate(expense_names, start=1):
        amt = _num(invoice.get(f"e{i}_amount"))
        rate = _num(invoice.get(f"e{i}_rate"))
        tot = amt * rate
        pdf.cell(exp_widths[0], 7, fa(name), border=1, align="R")
        pdf.cell(exp_widths[1], 7, f"{amt:,.2f}", border=1, align="C")
        pdf.cell(exp_widths[2], 7, f"{rate:,.0f}", border=1, align="C")
        pdf.cell(exp_widths[3], 7, f"{tot:,.0f}", border=1, align="C")
        pdf.ln()
        if i <= 7:
            e_taxable_total += tot
        else:
            e_avarez_total += tot
    pdf.ln(6)

    # ── خلاصه فاکتور دوم ──
    if pdf.has_bold:
        pdf.set_font("Nazanin", "B", size=11)
    else:
        pdf.set_font("Nazanin", size=11)
    pdf.set_text_color(20, 122, 111)
    pdf.cell(0, 7, fa("خلاصه فاکتور (فقط هزینه‌ها)"), align="R", new_x="LMARGIN", new_y="NEXT")

    pdf.set_font("Nazanin", size=10)
    pdf.set_text_color(40, 40, 40)
    exp_vat = e_taxable_total * 0.1
    exp_grand = e_taxable_total + exp_vat + e_avarez_total

    pdf.cell(0, 7, fa(f"جمع هزینه‌های مشمول (ردیف ۱ تا ۷): {e_taxable_total:,.0f} ریال"), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 7, fa(f"ارزش افزوده هزینه‌ها (۱۰٪): {exp_vat:,.0f} ریال"), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 7, fa(f"عوارض دامپزشکی (بدون ارزش افزوده): {e_avarez_total:,.0f} ریال"), align="R", new_x="LMARGIN", new_y="NEXT")
    if pdf.has_bold:
        pdf.set_font("Nazanin", "B", size=11)
    else:
        pdf.set_font("Nazanin", size=11)
    pdf.cell(0, 7, fa(f"جمع کل هزینه‌ها: {exp_grand:,.0f} ریال"), align="R", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)

    # ── توضیحات ──
    desc = invoice.get("description") or ""
    if desc.strip():
        if pdf.has_bold:
            pdf.set_font("Nazanin", "B", size=11)
        else:
            pdf.set_font("Nazanin", size=11)
        pdf.set_text_color(20, 122, 111)
        pdf.cell(0, 7, fa("توضیحات:"), align="R", new_x="LMARGIN", new_y="NEXT")

        pdf.set_font("Nazanin", size=9)
        pdf.set_text_color(40, 40, 40)

        lines = textwrap.wrap(desc, width=85)
        box_padding = 2
        line_height = 5.5
        box_h = (box_padding * 2) + (len(lines) * line_height)

        pdf.set_draw_color(220, 220, 220)
        pdf.rect(10, pdf.get_y(), 190, box_h)

        text_y = pdf.get_y() + box_padding
        for line in lines:
            pdf.set_xy(12, text_y)
            pdf.cell(186, line_height, fa(line), border=0, align="R")
            text_y += line_height

    pdf.output(output_path)
    return output_path


# ==========================================
# Excel
# ==========================================

def generate_daily_invoice_excel(form_date, output_path=None):
    from functions import get_daily_invoice
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter

    invoice = get_daily_invoice(form_date)
    if not invoice:
        raise ValueError(f"صورتحساب روز {form_date} یافت نشد.")

    os.makedirs(REPORTS_DIR, exist_ok=True)
    if output_path is None:
        safe_date = form_date.replace("/", "-")
        output_path = os.path.join(REPORTS_DIR, f"invoice_{safe_date}.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "صورتحساب کشتار روزانه"
    ws.sheet_view.rightToLeft = True

    # تنظیم عرض ستون‌ها
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 14
    ws.column_dimensions["G"].width = 20

    font_title = Font(name="B Nazanin", size=16, bold=True, color="147A6F")
    font_header = Font(name="B Nazanin", size=11, bold=True, color="FFFFFF")
    font_normal = Font(name="B Nazanin", size=11)
    font_bold = Font(name="B Nazanin", size=11, bold=True)

    fill_primary = PatternFill(start_color="147A6F", end_color="147A6F", fill_type="solid")
    fill_total = PatternFill(start_color="E0F2F1", end_color="E0F2F1", fill_type="solid")
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    border_thin = Border(
        left=Side(style="thin", color="999999"),
        right=Side(style="thin", color="999999"),
        top=Side(style="thin", color="999999"),
        bottom=Side(style="thin", color="999999"),
    )

    ws.merge_cells("A1:G1")
    c = ws["A1"]
    c.value = f"صورتحساب کشتار روزانه — {form_date}"
    c.font = font_title
    c.alignment = align_center
    ws.row_dimensions[1].height = 30

    current_row = 3

    # اطلاعات بالایی
    ws.cell(row=current_row, column=1, value="شماره صورتحساب:").font = font_bold
    ws.cell(row=current_row, column=2, value=invoice.get("invoice_number") or "").font = font_normal
    ws.cell(row=current_row, column=4, value="تاریخ:").font = font_bold
    ws.cell(row=current_row, column=5, value=invoice.get("invoice_date") or "").font = font_normal
    current_row += 2

    # ── جدول وزن‌کشی ──
    ws.cell(row=current_row, column=1, value="جدول وزن‌کشی").font = font_bold
    current_row += 1

    headers = ["ردیف", "وزن پر", "وزن خالی", "خالص", "تعداد", "فی", "مبلغ"]
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=current_row, column=i, value=h)
        cell.font = font_header
        cell.fill = fill_primary
        cell.alignment = align_center
        cell.border = border_thin
    current_row += 1

    total_weight = 0
    total_amount = 0
    for i in range(1, 6):
        full_val = _num(invoice.get(f"w{i}_full"))
        empty_val = _num(invoice.get(f"w{i}_empty"))
        net_val = max(0, full_val - empty_val)
        qty_val = _num(invoice.get(f"w{i}_qty"))
        rate_val = _num(invoice.get(f"w{i}_rate"))
        amt_val = net_val * (qty_val if qty_val > 0 else 1) * rate_val

        row_vals = [
            i,
            f"{full_val:,.2f}",
            f"{empty_val:,.2f}",
            f"{net_val:,.2f}",
            f"{qty_val:,.0f}",
            f"{rate_val:,.0f}",
            f"{amt_val:,.0f}",
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = font_normal
            cell.alignment = align_center
            cell.border = border_thin
        total_weight += net_val
        total_amount += amt_val
        current_row += 1

    row_vals = ["جمع", "", "", f"{total_weight:,.2f}", "", "", f"{total_amount:,.0f}"]
    for col_idx, val in enumerate(row_vals, start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=val)
        cell.font = font_bold
        cell.fill = fill_total
        cell.alignment = align_center
        cell.border = border_thin
    current_row += 1

    # ── کشتار و بسته‌بندی (چسبیده به جدول بالا) ──
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=7)
    cell = ws.cell(row=current_row, column=1, value="کشتار و بسته‌بندی")
    cell.font = font_header
    cell.fill = fill_primary
    cell.alignment = align_center
    cell.border = border_thin
    for c_idx in range(2, 8):
        ws.cell(row=current_row, column=c_idx).border = border_thin
        ws.cell(row=current_row, column=c_idx).fill = fill_primary
    current_row += 1

    for label, amount_key, rate_key, total_key in [
        ("کشتار مرغ گوشتی", "c1_amount", "c1_rate", "c1_total"),
        ("بسته بندی مرغ گوشتی", "c2_amount", "c2_rate", "c2_total"),
    ]:
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
        cell = ws.cell(row=current_row, column=1, value=label)
        cell.font = font_normal
        cell.alignment = align_right
        for c_idx in [1, 2]:
            ws.cell(row=current_row, column=c_idx).border = border_thin

        ws.cell(row=current_row, column=3, value=f"{_num(invoice.get(amount_key)):,.2f}").font = font_normal
        ws.cell(row=current_row, column=3).alignment = align_center
        ws.cell(row=current_row, column=3).border = border_thin

        ws.cell(row=current_row, column=4, value=f"{_num(invoice.get(rate_key)):,.0f}").font = font_normal
        ws.cell(row=current_row, column=4).alignment = align_center
        ws.cell(row=current_row, column=4).border = border_thin

        ws.merge_cells(start_row=current_row, start_column=5, end_row=current_row, end_column=7)
        cell = ws.cell(row=current_row, column=5, value=f"{_num(invoice.get(total_key)):,.0f}")
        cell.font = font_bold
        cell.alignment = align_center
        for c_idx in [5, 6, 7]:
            ws.cell(row=current_row, column=c_idx).border = border_thin

        current_row += 1

    current_row += 1

    # ── خلاصه فاکتور اول ──
    ws.cell(row=current_row, column=1, value="خلاصه فاکتور (کشتار و بسته‌بندی)").font = font_bold
    current_row += 1
    base = _num(invoice.get("c1_total")) + _num(invoice.get("c2_total"))
    ws.cell(row=current_row, column=1, value="مبلغ پایه (کشتار + بسته‌بندی):").font = font_normal
    ws.cell(row=current_row, column=2, value=f"{base:,.0f}").font = font_bold
    current_row += 1
    ws.cell(row=current_row, column=1, value="ارزش افزوده (۱۰٪):").font = font_normal
    ws.cell(row=current_row, column=2, value=f"{_num(invoice.get('vat_amount')):,.0f}").font = font_bold
    current_row += 1
    ws.cell(row=current_row, column=1, value="جمع کل فاکتور:").font = font_bold
    ws.cell(row=current_row, column=2, value=f"{_num(invoice.get('grand_total')):,.0f}").font = font_bold
    current_row += 2

    # ── هزینه‌ها ──
    ws.cell(row=current_row, column=1, value="هزینه‌ها").font = font_bold
    current_row += 1

    exp_headers = ["شرح هزینه", "وزن (کیلوگرم)", "فی (ریال)", "مبلغ (ریال)"]
    for i, h in enumerate(exp_headers, start=1):
        cell = ws.cell(row=current_row, column=i, value=h)
        cell.font = font_header
        cell.fill = fill_primary
        cell.alignment = align_center
        cell.border = border_thin
    current_row += 1

    expense_names = [
        "هزینه یخ مازاد داخل خودرو",
        "بسته بندی دل و کگر و سنگدان",
        "قطعه بندی مرغ",
        "پاکسازی دل و کگر و سنگدان",
        "پوست گیری پای مرغ",
        "صدور مجوز",
        "پالمپ",
        "عوارض دامپزشکی"
    ]

    e_taxable_total = 0.0
    e_avarez_total = 0.0
    for i, name in enumerate(expense_names, start=1):
        amt = _num(invoice.get(f"e{i}_amount"))
        rate = _num(invoice.get(f"e{i}_rate"))
        tot = amt * rate

        ws.cell(row=current_row, column=1, value=name).font = font_normal
        ws.cell(row=current_row, column=2, value=f"{amt:,.2f}").font = font_normal
        ws.cell(row=current_row, column=3, value=f"{rate:,.0f}").font = font_normal
        ws.cell(row=current_row, column=4, value=f"{tot:,.0f}").font = font_bold
        current_row += 1

        if i <= 7:
            e_taxable_total += tot
        else:
            e_avarez_total += tot

    current_row += 1

    # ── خلاصه فاکتور دوم (فقط هزینه‌ها) ──
    ws.cell(row=current_row, column=1, value="خلاصه فاکتور (فقط هزینه‌ها)").font = font_bold
    current_row += 1

    exp_vat = e_taxable_total * 0.1
    exp_grand = e_taxable_total + exp_vat + e_avarez_total

    ws.cell(row=current_row, column=1, value="جمع هزینه‌های مشمول (ردیف ۱ تا ۷):").font = font_normal
    ws.cell(row=current_row, column=2, value=f"{e_taxable_total:,.0f}").font = font_bold
    current_row += 1
    ws.cell(row=current_row, column=1, value="ارزش افزوده هزینه‌ها (۱۰٪):").font = font_normal
    ws.cell(row=current_row, column=2, value=f"{exp_vat:,.0f}").font = font_bold
    current_row += 1
    ws.cell(row=current_row, column=1, value="عوارض دامپزشکی (بدون ارزش افزوده):").font = font_normal
    ws.cell(row=current_row, column=2, value=f"{e_avarez_total:,.0f}").font = font_bold
    current_row += 1
    ws.cell(row=current_row, column=1, value="جمع کل هزینه‌ها:").font = font_bold
    ws.cell(row=current_row, column=2, value=f"{exp_grand:,.0f}").font = font_bold

    wb.save(output_path)
    return output_path
