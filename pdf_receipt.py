# -*- coding: utf-8 -*-
"""
تولید PDF و Excel برای قبض باسکول — نسخه نهایی
"""

import os
from datetime import datetime
from fpdf import FPDF

import arabic_reshaper
from bidi.algorithm import get_display

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
FONT_REG_PATH = os.path.join(BASE_DIR, "assets", "fonts", "BNazanin.ttf")
FONT_BOLD_PATH = os.path.join(BASE_DIR, "assets", "fonts", "BNazaninBold.ttf")
REPORTS_DIR = os.path.join(BASE_DIR, "reports")


def fa(text):
    if not text:
        return ""
    reshaped = arabic_reshaper.reshape(str(text))
    return get_display(reshaped)


def _try_float(v):
    if not v:
        return None
    try:
        return float(v)
    except (ValueError, TypeError):
        return None


def parse_extra(description):
    if not description:
        return {
            "full": None, "empty": None,
            "code": "", "name": "", "title": "",
            "text": "", "full_date": "", "full_time": "",
            "empty_date": "", "empty_time": "",
            "scale_name": "", "weighing_cost": "",
        }
    parts = {}
    text_parts = []
    for token in description.split("|"):
        if ":" in token:
            k, v = token.split(":", 1)
            k = k.strip().upper()
            if k in ("FULL", "EMPTY", "CODE", "NAME", "TITLE",
                     "FULL_DATE", "FULL_TIME", "EMPTY_DATE", "EMPTY_TIME",
                     "SCALE", "COST"):
                parts[k] = v.strip()
                continue
        text_parts.append(token)

    if parts:
        return {
            "full": _try_float(parts.get("FULL")),
            "empty": _try_float(parts.get("EMPTY")),
            "code": parts.get("CODE", ""),
            "name": parts.get("NAME", ""),
            "title": parts.get("TITLE", ""),
            "full_date": parts.get("FULL_DATE", ""),
            "full_time": parts.get("FULL_TIME", ""),
            "empty_date": parts.get("EMPTY_DATE", ""),
            "empty_time": parts.get("EMPTY_TIME", ""),
            "scale_name": parts.get("SCALE", ""),
            "weighing_cost": parts.get("COST", ""),
            "text": " | ".join(text_parts).strip(),
        }
    return {
        "full": None, "empty": None,
        "code": "", "name": "", "title": "",
        "full_date": "", "full_time": "", "empty_date": "", "empty_time": "",
        "scale_name": "", "weighing_cost": "",
        "text": description,
    }


# رنگ‌ها برای PDF
DARK_TEXT = (40, 40, 40)
PRIMARY = (20, 122, 111)
PRIMARY_LIGHT = (235, 245, 245)
GRAY = (100, 100, 100)
LIGHT_GRAY = (220, 220, 220)
WHITE = (255, 255, 255)


class ReceiptPDF(FPDF):
    def __init__(self, scale_name="باسکول مرکزی"):
        super().__init__(orientation="P", unit="mm", format="A5")
        self.set_auto_page_break(auto=False)
        self.set_margins(10, 10, 10)
        self.scale_name = scale_name

        if not os.path.exists(FONT_REG_PATH):
            raise FileNotFoundError(f"فونت پیدا نشد: {FONT_REG_PATH}")

        self.add_font("Nazanin", "", FONT_REG_PATH)
        if os.path.exists(FONT_BOLD_PATH):
            self.add_font("Nazanin", "B", FONT_BOLD_PATH)
            self.has_bold = True
        else:
            self.has_bold = False

    def header(self):
        if self.has_bold:
            self.set_font("Nazanin", "B", size=14)
        else:
            self.set_font("Nazanin", size=14)
        self.set_text_color(*PRIMARY)
        self.cell(0, 8, fa(f"باسکول: {self.scale_name}"), align="C",
                  new_x="LMARGIN", new_y="NEXT")

        self.ln(2)
        self.set_draw_color(*PRIMARY)
        self.set_line_width(0.5)
        self.line(10, self.get_y(), 138, self.get_y())
        self.ln(3)

    def footer(self):
        self.set_y(-10)
        self.set_font("Nazanin", size=8)
        self.set_text_color(*GRAY)
        self.cell(0, 6, fa(f"تاریخ چاپ: {datetime.now().strftime('%Y/%m/%d %H:%M')}"), align="C")


def generate_receipt_pdf(receipt_data, slaughter_info, output_path=None, scale_name=None):
    os.makedirs(REPORTS_DIR, exist_ok=True)

    extra = parse_extra(receipt_data.get("description") or "")
    final_scale_name = scale_name or extra.get("scale_name") or "باسکول مرکزی"

    if output_path is None:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        receipt_num = receipt_data.get("receipt_number") or "unknown"
        filename = f"receipt_{receipt_num}_{ts}.pdf"
        output_path = os.path.join(REPORTS_DIR, filename)

    pdf = ReceiptPDF(scale_name=final_scale_name)
    pdf.add_page()

    receipt_num = receipt_data.get("receipt_number") or "—"
    date_str = receipt_data.get("receipt_date") or "—"
    car_plate = receipt_data.get("car_plate") or "—"
    driver_name = receipt_data.get("driver_name") or "—"
    buyer_name = receipt_data.get("buyer_name") or "—"
    slaughter_num = f"#{slaughter_info['slaughter_number']}"

    y = pdf.get_y()

    # ── اطلاعات بالایی ──
    pdf.set_font("Nazanin", size=9)
    pdf.set_text_color(*DARK_TEXT)

    x_right = 96
    w_right = 42
    x_mid = 54
    w_mid = 42
    x_left = 12
    w_left = 42

    pdf.set_xy(x_right, y)
    pdf.cell(w_right, 6, fa(f"شماره کشتار: {slaughter_num}"), align="R")
    pdf.set_xy(x_mid, y)
    pdf.cell(w_mid, 6, fa(f"شماره قبض: {receipt_num}"), align="R")
    y += 6

    pdf.set_xy(x_right, y)
    pdf.cell(w_right, 6, fa(f"مشتری: {buyer_name}"), align="R")
    pdf.set_xy(x_mid, y)
    pdf.cell(w_mid, 6, fa(f"راننده: {driver_name}"), align="R")
    pdf.set_xy(x_left, y)
    pdf.cell(w_left, 6, fa(f"پلاک ماشین: {car_plate}"), align="R")
    y += 8

    # ── جدول توزین ──
    w1 = 24
    w2 = 32
    w3 = 32
    w4 = 40
    total_w = w1 + w2 + w3 + w4
    right_x = 10 + total_w

    if pdf.has_bold:
        pdf.set_font("Nazanin", "B", size=9)
    else:
        pdf.set_font("Nazanin", size=9)
    pdf.set_text_color(*WHITE)
    pdf.set_fill_color(*PRIMARY)

    headers = ["عنوان", "ساعت", "تاریخ", "وزن (کیلوگرم)"]
    widths = [w1, w2, w3, w4]

    current_x = right_x
    pdf.set_xy(current_x, y)
    for h, w in zip(headers, widths):
        current_x -= w
        pdf.set_xy(current_x, y)
        pdf.cell(w, 7, fa(h), border=1, align="C", fill=True)
    pdf.ln()
    y += 7

    # ردیف ۱: خالی
    pdf.set_font("Nazanin", size=9)
    pdf.set_text_color(*DARK_TEXT)
    current_x = right_x

    current_x -= w1
    pdf.set_xy(current_x, y)
    pdf.cell(w1, 7, fa("خالی"), border=1, align="C")
    current_x -= w2
    pdf.set_xy(current_x, y)
    pdf.cell(w2, 7, fa(extra['empty_time'] or "—"), border=1, align="C")
    current_x -= w3
    pdf.set_xy(current_x, y)
    pdf.cell(w3, 7, fa(extra['empty_date'] or date_str), border=1, align="C")
    current_x -= w4
    pdf.set_xy(current_x, y)
    pdf.cell(w4, 7, f"{extra['empty']:,.2f}" if extra['empty'] else "—", border=1, align="C")
    pdf.ln()
    y += 7

    # ردیف ۲: پر
    pdf.set_font("Nazanin", size=9)
    current_x = right_x

    current_x -= w1
    pdf.set_xy(current_x, y)
    pdf.cell(w1, 7, fa("پر"), border=1, align="C")
    current_x -= w2
    pdf.set_xy(current_x, y)
    pdf.cell(w2, 7, fa(extra['full_time'] or "—"), border=1, align="C")
    current_x -= w3
    pdf.set_xy(current_x, y)
    pdf.cell(w3, 7, fa(extra['full_date'] or date_str), border=1, align="C")
    current_x -= w4
    pdf.set_xy(current_x, y)
    pdf.cell(w4, 7, f"{extra['full']:,.2f}" if extra['full'] else "—", border=1, align="C")
    pdf.ln()
    y += 7

    # ردیف ۳: نام کالا، کد کالا، خالص
    pdf.set_font("Nazanin", size=9)
    pdf.set_text_color(*DARK_TEXT)
    current_x = right_x

    current_x -= (w1 + w2)
    pdf.set_xy(current_x, y)
    pdf.cell(w1 + w2, 7, fa(f"نام کالا: {extra['name'] or '—'}"), border=1, align="C")
    current_x -= w3
    pdf.set_xy(current_x, y)
    pdf.cell(w3, 7, fa(f"کد کالا: {extra['code'] or '—'}"), border=1, align="C")
    current_x -= w4
    pdf.set_xy(current_x, y)
    net_val = f"{receipt_data.get('net_weight', 0):,.2f}"
    pdf.cell(w4, 7, fa(f"{net_val} :خالص"), border=1, align="C")
    pdf.ln()
    y += 7

    # ردیف ۴: هزینه توزین
    pdf.set_font("Nazanin", size=9)
    pdf.set_text_color(*DARK_TEXT)
    current_x = right_x - total_w
    pdf.set_xy(current_x, y)
    cost_val = extra.get("weighing_cost") or "—"
    pdf.cell(total_w, 7, fa(f"هزینه توزین: {cost_val}"), border=1, align="C")
    pdf.ln()
    y += 10

    # توضیحات
    desc = extra["text"]
    if desc.strip():
        if pdf.has_bold:
            pdf.set_font("Nazanin", "B", size=10)
        else:
            pdf.set_font("Nazanin", size=10)
        pdf.set_text_color(*PRIMARY)
        pdf.set_xy(10, y)
        pdf.cell(128, 6, fa("توضیحات:"), align="R")
        y += 6

        pdf.set_font("Nazanin", size=9)
        pdf.set_text_color(*DARK_TEXT)

        max_chars = 75
        words = desc.split(" ")
        lines = []
        current_line = ""
        for word in words:
            if len(current_line) + len(word) + 1 > max_chars:
                if current_line:
                    lines.append(current_line.strip())
                current_line = word + " "
            else:
                current_line += word + " "

        if current_line.strip():
            lines.append(current_line.strip())

        box_padding = 2
        line_height = 5.5
        box_h = (box_padding * 2) + (len(lines) * line_height)

        pdf.set_draw_color(*LIGHT_GRAY)
        pdf.rect(10, y, 128, box_h)

        text_y = y + box_padding
        for line in lines:
            pdf.set_xy(12, text_y)
            pdf.cell(124, line_height, fa(line), border=0, align="R")
            text_y += line_height

    pdf.output(output_path)
    return output_path


# ==========================================
# تولید Excel (xlsx)
# ==========================================

def generate_receipt_excel(receipt_data, slaughter_info, output_path=None, scale_name=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter

    os.makedirs(REPORTS_DIR, exist_ok=True)

    extra = parse_extra(receipt_data.get("description") or "")
    final_scale_name = scale_name or extra.get("scale_name") or "باسکول مرکزی"

    if output_path is None:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        receipt_num = receipt_data.get("receipt_number") or "unknown"
        filename = f"receipt_{receipt_num}_{ts}.xlsx"
        output_path = os.path.join(REPORTS_DIR, filename)

    wb = Workbook()
    ws = wb.active
    ws.title = "قبض باسکول"
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False

    # استایل‌ها
    font_title = Font(name="B Nazanin", size=16, bold=True, color="147A6F")
    font_normal = Font(name="B Nazanin", size=11)
    font_bold = Font(name="B Nazanin", size=11, bold=True)
    
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    border_thin = Border(
        left=Side(style="thin", color="000000"),
        right=Side(style="thin", color="000000"),
        top=Side(style="thin", color="000000"),
        bottom=Side(style="thin", color="000000"),
    )
    
    fill_primary = PatternFill(start_color="147A6F", end_color="147A6F", fill_type="solid")
    font_white = Font(name="B Nazanin", size=11, bold=True, color="FFFFFF")

    # تنظیم عرض ستون‌ها (طبق فرمت PDF)
    col_widths = {"A": 18, "B": 18, "C": 18, "D": 18, "E": 18, "F": 20}
    for col, w in col_widths.items():
        ws.column_dimensions[col].width = w

    # --- عنوان (باسکول) ---
    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = f"باسکول: {final_scale_name}"
    c.font = font_title
    c.alignment = align_center
    ws.row_dimensions[1].height = 30

    # خط جداکننده (شبیه‌سازی با کادر پایین)
    ws.merge_cells("A2:F2")
    c = ws["A2"]
    c.border = Border(bottom=Side(style="medium", color="147A6F"))
    
    # --- اطلاعات بالایی ---
    # سطر ۱
    ws.merge_cells("A4:C4")
    ws["A4"] = f"شماره کشتار: #{slaughter_info['slaughter_number']}"
    ws["A4"].font = font_normal
    ws["A4"].alignment = align_right
    
    ws.merge_cells("D4:F4")
    ws["D4"] = f"شماره قبض: {receipt_data.get('receipt_number') or '—'}"
    ws["D4"].font = font_normal
    ws["D4"].alignment = align_right

    # سطر ۲
    ws.merge_cells("A5:B5")
    ws["A5"] = f"مشتری: {receipt_data.get('buyer_name') or '—'}"
    ws["A5"].font = font_normal
    ws["A5"].alignment = align_right

    ws.merge_cells("C5:D5")
    ws["C5"] = f"راننده: {receipt_data.get('driver_name') or '—'}"
    ws["C5"].font = font_normal
    ws["C5"].alignment = align_right

    ws.merge_cells("E5:F5")
    ws["E5"] = f"پلاک ماشین: {receipt_data.get('car_plate') or '—'}"
    ws["E5"].font = font_normal
    ws["E5"].alignment = align_right
    
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 20

    # --- جدول توزین ---
    start_row = 7
    
    # هدر جدول
    headers = ["عنوان", "ساعت", "تاریخ", "وزن (کیلوگرم)"]
    # ادغام برای هماهنگی با ستون‌های A تا F:
    # عنوان: A, ساعت: B-C, تاریخ: D-E, وزن: F
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=1)
    ws.merge_cells(start_row=start_row, start_column=2, end_row=start_row, end_column=3)
    ws.merge_cells(start_row=start_row, start_column=4, end_row=start_row, end_column=5)
    ws.merge_cells(start_row=start_row, start_column=6, end_row=start_row, end_column=6)
    
    col_map = [1, 2, 4, 6]
    
    for i, h in enumerate(headers):
        cell = ws.cell(row=start_row, column=col_map[i], value=h)
        cell.font = font_white
        cell.fill = fill_primary
        cell.alignment = align_center
        # اعمال حاشیه به تمام سلول‌های ادغام‌شده
        for c_idx in range(col_map[i], col_map[i] + (2 if i in [1, 2] else 1)):
            ws.cell(row=start_row, column=c_idx).border = border_thin
            ws.cell(row=start_row, column=c_idx).fill = fill_primary
            
    ws.row_dimensions[start_row].height = 25

    # ردیف ۱: خالی
    row1 = start_row + 1
    ws.merge_cells(start_row=row1, start_column=2, end_row=row1, end_column=3)
    ws.merge_cells(start_row=row1, start_column=4, end_row=row1, end_column=5)
    
    vals1 = ["خالی", extra['empty_time'] or "—", extra['empty_date'] or "—", f"{extra['empty']:,.2f}" if extra['empty'] else "—"]
    for i, v in enumerate(vals1):
        cell = ws.cell(row=row1, column=col_map[i], value=v)
        cell.font = font_normal
        cell.alignment = align_center
        for c_idx in range(col_map[i], col_map[i] + (2 if i in [1, 2] else 1)):
            ws.cell(row=row1, column=c_idx).border = border_thin
    ws.row_dimensions[row1].height = 22

    # ردیف ۲: پر
    row2 = start_row + 2
    ws.merge_cells(start_row=row2, start_column=2, end_row=row2, end_column=3)
    ws.merge_cells(start_row=row2, start_column=4, end_row=row2, end_column=5)
    
    vals2 = ["پر", extra['full_time'] or "—", extra['full_date'] or "—", f"{extra['full']:,.2f}" if extra['full'] else "—"]
    for i, v in enumerate(vals2):
        cell = ws.cell(row=row2, column=col_map[i], value=v)
        cell.font = font_normal
        cell.alignment = align_center
        for c_idx in range(col_map[i], col_map[i] + (2 if i in [1, 2] else 1)):
            ws.cell(row=row2, column=c_idx).border = border_thin
    ws.row_dimensions[row2].height = 22

    # ردیف ۳: نام کالا، کد کالا، خالص
    row3 = start_row + 3
    ws.merge_cells(start_row=row3, start_column=1, end_row=row3, end_column=3)
    ws.merge_cells(start_row=row3, start_column=4, end_row=row3, end_column=5)
    
    net_val = f"{receipt_data.get('net_weight', 0):,.2f}"
    vals3 = [f"نام کالا: {extra['name'] or '—'}", f"کد کالا: {extra['code'] or '—'}", "", f"{net_val} :خالص"]
    for i, v in enumerate(vals3):
        cell = ws.cell(row=row3, column=col_map[i], value=v)
        cell.font = font_normal
        cell.alignment = align_center
        for c_idx in range(col_map[i], col_map[i] + (2 if i in [1, 2] else 1)):
            ws.cell(row=row3, column=c_idx).border = border_thin
    ws.row_dimensions[row3].height = 22

    # ردیف ۴: هزینه توزین
    row4 = start_row + 4
    ws.merge_cells(start_row=row4, start_column=1, end_row=row4, end_column=6)
    cell = ws.cell(row=row4, column=1, value=f"هزینه توزین: {extra.get('weighing_cost') or '—'}")
    cell.font = font_normal
    cell.alignment = align_center
    for c_idx in range(1, 7):
        ws.cell(row=row4, column=c_idx).border = border_thin
    ws.row_dimensions[row4].height = 22

    # --- توضیحات ---
    desc = extra["text"]
    if desc.strip():
        start_desc = row4 + 2
        ws.merge_cells(start_row=start_desc, start_column=1, end_row=start_desc, end_column=6)
        cell = ws.cell(row=start_desc, column=1, value="توضیحات:")
        cell.font = font_bold
        cell.alignment = align_right
        
        start_box = start_desc + 1
        ws.merge_cells(start_row=start_box, start_column=1, end_row=start_box, end_column=6)
        cell = ws.cell(row=start_box, column=1, value=desc)
        cell.font = font_normal
        cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        
        for c_idx in range(1, 7):
            ws.cell(row=start_box, column=c_idx).border = Border(
                left=Side(style="medium", color="CCCCCC"),
                right=Side(style="medium", color="CCCCCC"),
                top=Side(style="medium", color="CCCCCC"),
                bottom=Side(style="medium", color="CCCCCC"),
            )
        
        # تخمین ارتفاع برای متن
        lines_count = max(1, len(desc) // 75 + 1)
        ws.row_dimensions[start_box].height = lines_count * 20

    wb.save(output_path)
    return output_path
