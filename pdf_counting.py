# -*- coding: utf-8 -*-
"""
تولید PDF و Excel فرم آماری شمارش مرغ از دیتابیس
"""

import os
from datetime import datetime
from fpdf import FPDF

import arabic_reshaper
from bidi.algorithm import get_display

try:
    import jdatetime
    HAS_JDATETIME = True
except ImportError:
    HAS_JDATETIME = False


BASE_DIR = os.path.dirname(os.path.abspath(__file__))

FONT_REG_PATH = os.path.join(BASE_DIR, "assets", "fonts", "BNazanin.ttf")
FONT_BOLD_PATH = os.path.join(BASE_DIR, "assets", "fonts", "BNazaninBold.ttf")
LOGO_PATH = os.path.join(BASE_DIR, "assets", "images", "logo.png")

REPORTS_DIR = os.path.join(BASE_DIR, "reports")


def fa(text):
    if not text:
        return ""
    reshaped = arabic_reshaper.reshape(str(text))
    return get_display(reshaped)


def safe_int(value, default=0):
    try:
        return int(value) if value else default
    except (TypeError, ValueError):
        return default


def truncate_text(text, max_chars=35):
    if text is None or text == "":
        return ""
    text = str(text)
    if len(text) <= max_chars:
        return text
    return text[:max_chars - 3] + "..."


BLACK = (0, 0, 0)
WHITE = (255, 255, 255)
RED = (200, 0, 0)
GRAY = (220, 220, 220)

LEFT_EDGE = 15
RIGHT_EDGE = 195
TOTAL_WIDTH = RIGHT_EDGE - LEFT_EDGE


# ==========================================
# محاسبات
# ==========================================

def calc_packed_expected(hook_count, waste, pieces):
    hc = safe_int(hook_count)
    w = safe_int(waste)
    pcs = safe_int(pieces)
    expected = hc - w - pcs
    return expected if expected >= 0 else 0


def calc_discrepancy(packed_entered, packed_expected):
    return packed_entered - packed_expected


def calc_loaded_on_truck(hook_count, deaths):
    return safe_int(hook_count) + safe_int(deaths)


def format_discrepancy(diff):
    if diff == 0:
        return None
    if diff > 0:
        return str(diff)
    return f"{abs(diff)}-"


# ==========================================
# کلاس PDF
# ==========================================

class CountingFormPDF(FPDF):

    def __init__(self):
        super().__init__(orientation="P", unit="mm", format="A4")
        self.set_auto_page_break(auto=False)
        self.set_margins(15, 15, 15)

        if not os.path.exists(FONT_REG_PATH):
            raise FileNotFoundError(f"فونت پیدا نشد: {FONT_REG_PATH}")

        self.add_font("Nazanin", "", FONT_REG_PATH)

        if os.path.exists(FONT_BOLD_PATH):
            self.add_font("Nazanin", "B", FONT_BOLD_PATH)
            self.has_bold = True
        else:
            self.has_bold = False

    def draw_header(self, form_date, day_name=""):
        if os.path.exists(LOGO_PATH):
            try:
                self.image(LOGO_PATH, x=15, y=10, w=20)
            except Exception:
                pass

        self.set_y(15)
        if self.has_bold:
            self.set_font("Nazanin", "B", size=18)
        else:
            self.set_font("Nazanin", size=18)

        self.set_text_color(*BLACK)
        self.cell(180, 12, fa("فرم آماری شمارش مرغ"), align="C")

        self.set_y(35)
        self.set_font("Nazanin", size=12)
        self.set_text_color(*BLACK)

        date_text = f"تاریخ: {form_date}"
        if day_name:
            date_text += f"  ({day_name})"

        self.set_xy(120, 35)
        self.cell(75, 7, fa(date_text), align="R")

        self.set_y(48)

    def draw_table(self, rows, y):
        padding_top = 1.5
        padding_bottom = 1.5
        line_gap = 0.3
        text_height = 3.5
        header_font_size = 8

        header_h = (padding_top + text_height + line_gap
                    + text_height + padding_bottom)

        row_height = 8
        row_font_size = 9

        unit = TOTAL_WIDTH / 24

        widths = [
            5.0 * unit,
            3.5 * unit,
            3.0 * unit,
            2.5 * unit,
            2.5 * unit,
            3.5 * unit,
            4.0 * unit,
        ]

        headers = [
            ("توضیحات", ""),
            ("بسته‌بندی", "(قطعه)"),
            ("قطعه‌بندی", "(قطعه)"),
            ("تلفات", "(قطعه)"),
            ("ضایعات", "(قطعه)"),
            ("پای قلاب", "(قطعه)"),
            ("کشتار", ""),
        ]

        self.set_text_color(*BLACK)
        self.set_draw_color(*BLACK)

        start_x = LEFT_EDGE
        start_y = y
        current_x = start_x

        for (title, sub), w in zip(headers, widths):
            self.set_xy(current_x, start_y + padding_top)
            self.set_font("Nazanin", size=header_font_size)
            self.cell(w, text_height, fa(title),
                      border=0, align="C")

            self.set_xy(current_x,
                        start_y + padding_top + text_height + line_gap)
            self.set_font("Nazanin", size=header_font_size)
            self.cell(w, text_height, fa(sub) if sub else "",
                      border=0, align="C")

            current_x += w

        self.rect(start_x, start_y, TOTAL_WIDTH, header_h)

        x_pos = start_x
        for w in widths:
            x_pos += w
            self.line(x_pos, start_y, x_pos, start_y + header_h)

        self.set_xy(LEFT_EDGE, start_y + header_h)

        for row in rows:
            self.set_text_color(*BLACK)
            self.set_font("Nazanin", size=row_font_size)

            notes_display = truncate_text(row["notes"], max_chars=35)

            self.cell(widths[0], row_height,
                      fa(notes_display) if notes_display else "",
                      border=1, align="R")
            self.cell(widths[1], row_height,
                      f"{row['packed']:,}" if row['packed'] else "",
                      border=1, align="C")
            self.cell(widths[2], row_height,
                      f"{row['pieces']:,}" if row['pieces'] else "",
                      border=1, align="C")
            self.cell(widths[3], row_height,
                      f"{row['deaths']:,}" if row['deaths'] else "",
                      border=1, align="C")
            self.cell(widths[4], row_height,
                      f"{row['waste']:,}" if row['waste'] else "",
                      border=1, align="C")
            self.cell(widths[5], row_height,
                      f"{row['hook_count']:,}" if row['hook_count'] else "",
                      border=1, align="C")
            self.cell(widths[6], row_height,
                      str(row["row_num"]),
                      border=1, align="C")

            self.ln()

        return self.get_y()

    def draw_info_lines(self, rows, y):
        items = []

        for row in rows:
            packed = safe_int(row.get("packed", 0))
            expected = calc_packed_expected(
                row.get("hook_count", 0),
                row.get("waste", 0),
                row.get("pieces", 0),
            )
            diff = calc_discrepancy(packed, expected)

            hook = safe_int(row.get("hook_count", 0))
            deaths = safe_int(row.get("deaths", 0))
            loaded = calc_loaded_on_truck(hook, deaths)

            items.append({
                "row": row["row_num"],
                "diff_str": format_discrepancy(diff) if diff != 0 else None,
                "hook": hook,
                "waste": safe_int(row.get("waste", 0)),
                "pieces": safe_int(row.get("pieces", 0)),
                "loaded": loaded,
            })

        if not items:
            return y

        y += 8

        self.set_font("Nazanin", size=11)

        for item in items:
            if item["diff_str"] is not None:
                self.set_text_color(*RED)
                text = (
                    f"ردیف {item['row']} — "
                    f"مغایرت {item['diff_str']} = "
                    f"{item['pieces']:,} - "
                    f"{item['waste']:,} - "
                    f"{item['hook']:,}"
                )
                self.set_xy(LEFT_EDGE, y)
                self.cell(TOTAL_WIDTH, 7, fa(text), align="R")
                y += 7

            self.set_text_color(*BLACK)
            text = (
                f"ردیف {item['row']} — "
                f"بارگیری‌شده داخل ماشین: {item['loaded']:,}"
            )
            self.set_xy(LEFT_EDGE, y)
            self.cell(TOTAL_WIDTH, 7, fa(text), align="R")
            y += 7

        return y

    def draw_notes(self, rows, y):
        items = []

        for row in rows:
            notes = row.get("notes", "")
            if notes and len(str(notes)) > 35:
                items.append({
                    "row": row["row_num"],
                    "notes": str(notes),
                })

        if not items:
            return y

        y += 5

        # عنوان
        if self.has_bold:
            self.set_font("Nazanin", "B", size=12)
        else:
            self.set_font("Nazanin", size=12)

        self.set_text_color(*BLACK)
        self.set_xy(LEFT_EDGE, y)
        self.cell(TOTAL_WIDTH, 8, fa("توضیحات کامل:"), align="R")

        y += 9

        self.set_font("Nazanin", size=11)
        self.set_text_color(*BLACK)

        box_padding = 3
        line_height = 7
        max_chars_per_line = 95

        for item in items:
            full_text = f"کشتار {item['row']}: {item['notes']}"

            # ── ۱. شکستن دستی متن به خطوط (از بالا به پایین) ──
            lines = []
            words = full_text.split(" ")
            current_line = ""
            for word in words:
                if len(current_line) + len(word) + 1 > max_chars_per_line:
                    if current_line:
                        lines.append(current_line.strip())
                    current_line = word + " "
                else:
                    current_line += word + " "
            if current_line.strip():
                lines.append(current_line.strip())

            # ── ۲. محاسبه ارتفاع کادر ──
            box_h = (box_padding * 2) + (len(lines) * line_height)

            # ── ۳. رسم کادر با رنگ خاکستری ──
            self.set_draw_color(*GRAY)
            self.rect(LEFT_EDGE, y, TOTAL_WIDTH, box_h)

            # ── ۴. نوشتن خطوط از بالا به پایین ──
            text_y = y + box_padding
            for line in lines:
                # هر خط را در مختصات y مشخص می‌نویسیم
                self.set_xy(LEFT_EDGE + box_padding, text_y)
                self.cell(TOTAL_WIDTH - (box_padding * 2), line_height,
                          fa(line), border=0, align="R")
                text_y += line_height

            # رفتن به کادر بعدی
            y += box_h + 2

        # بازگرداندن رنگ خط به مشکی برای بخش‌های بعدی
        self.set_draw_color(*BLACK)

        return y


# ==========================================
# تابع تولید PDF
# ==========================================

def generate_counting_pdf(form_date, output_path=None):
    from functions import get_counting_rows, get_counting_row

    rows = get_counting_rows(form_date)

    if not rows:
        raise ValueError(f"هیچ ردیفی برای تاریخ {form_date} یافت نشد.")

    day_name = ""
    first = get_counting_row(rows[0]["row_id"])
    if first:
        day_name = first.get("day_name", "")

    pdf = CountingFormPDF()
    pdf.add_page()

    pdf.draw_header(form_date=form_date, day_name=day_name)

    y = pdf.get_y() + 5
    y = pdf.draw_table(rows, y)
    y = pdf.draw_info_lines(rows, y)
    y = pdf.draw_notes(rows, y)

    if output_path is None:
        os.makedirs(REPORTS_DIR, exist_ok=True)
        safe_date = form_date.replace("/", "-")
        ts = datetime.now().strftime("%H%M%S")
        filename = f"counting_{safe_date}_{ts}.pdf"
        output_path = os.path.join(REPORTS_DIR, filename)

    pdf.output(output_path)
    return output_path


# ==========================================
# تولید Excel (xlsx) — با openpyxl
# ==========================================

def generate_counting_excel(form_date, output_path=None):
    """
    تولید فایل Excel (xlsx) از فرم ۱ (شمارش) با استایل‌دهی مناسب.
    """
    from functions import get_counting_rows, get_counting_row

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ImportError(
            "کتابخانه openpyxl نصب نیست. لطفاً با دستور 'pip install openpyxl' نصب کنید."
        )

    rows = get_counting_rows(form_date)

    if not rows:
        raise ValueError(f"هیچ ردیفی برای تاریخ {form_date} یافت نشد.")

    day_name = ""
    first = get_counting_row(rows[0]["row_id"])
    if first:
        day_name = first.get("day_name", "")

    if output_path is None:
        os.makedirs(REPORTS_DIR, exist_ok=True)
        safe_date = form_date.replace("/", "-")
        ts = datetime.now().strftime("%H%M%S")
        filename = f"counting_{safe_date}_{ts}.xlsx"
        output_path = os.path.join(REPORTS_DIR, filename)

    wb = Workbook()
    ws = wb.active
    ws.title = "شمارش"
    ws.sheet_view.rightToLeft = True

    header_font = Font(name="B Nazanin", size=12, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1B9E8F", end_color="1B9E8F", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    title_font = Font(name="B Nazanin", size=16, bold=True, color="1B9E8F")
    subtitle_font = Font(name="B Nazanin", size=11, color="333333")

    cell_font = Font(name="B Nazanin", size=11, color="000000")
    cell_align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    total_font = Font(name="B Nazanin", size=11, bold=True, color="000000")
    total_fill = PatternFill(start_color="E0F2F1", end_color="E0F2F1", fill_type="solid")

    red_font = Font(name="B Nazanin", size=11, bold=True, color="C00000")
    green_font = Font(name="B Nazanin", size=11, bold=True, color="2E7D32")

    thin_border = Border(
        left=Side(style="thin", color="999999"),
        right=Side(style="thin", color="999999"),
        top=Side(style="thin", color="999999"),
        bottom=Side(style="thin", color="999999"),
    )

    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = "فرم آماری شمارش مرغ"
    c.font = title_font
    c.alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:F2")
    c = ws["A2"]
    date_text = f"تاریخ: {form_date}"
    if day_name:
        date_text += f"  ({day_name})"
    c.value = date_text
    c.font = subtitle_font
    c.alignment = Alignment(horizontal="center", vertical="center")

    start_row = 4

    headers = [
        "کشتار", "پای قلاب", "ضایعات", "تلفات", "قطعه‌بندی", "بسته‌بندی",
        "مغایرت", "بارگیری‌شده در ماشین", "توضیحات"
    ]

    col_widths = [8, 12, 10, 10, 12, 14, 12, 22, 50]

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

        ws.column_dimensions[get_column_letter(col_idx)].width = width

    current_row = start_row + 1

    for row in rows:
        hc = safe_int(row.get("hook_count", 0))
        w = safe_int(row.get("waste", 0))
        d = safe_int(row.get("deaths", 0))
        pcs = safe_int(row.get("pieces", 0))
        p = safe_int(row.get("packed", 0))

        expected = hc - w - pcs
        diff = p - expected
        diff_str = ""
        if diff != 0:
            diff_str = str(diff) if diff > 0 else f"{abs(diff)}-"

        loaded = hc + d

        values = [
            row["row_num"],
            hc,
            w,
            d,
            pcs,
            p,
            diff_str,
            loaded,
            row.get("notes", ""),
        ]

        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=value)
            cell.font = cell_font
            cell.alignment = cell_align_center if col_idx != 9 else cell_align_right
            cell.border = thin_border

            if col_idx == 7 and diff_str:
                cell.font = red_font

        current_row += 1

    ws.cell(row=current_row, column=1, value="جمع").font = total_font
    ws.cell(row=current_row, column=1).fill = total_fill
    ws.cell(row=current_row, column=1).alignment = cell_align_center
    ws.cell(row=current_row, column=1).border = thin_border

    for col_idx in range(2, 9):
        cell = ws.cell(row=current_row, column=col_idx, value="")
        cell.font = total_font
        cell.fill = total_fill
        cell.alignment = cell_align_center
        cell.border = thin_border

    numeric_cols = {2: 0, 3: 0, 4: 0, 5: 0, 6: 0, 8: 0}
    for r in range(start_row + 1, current_row):
        for col_idx in numeric_cols:
            val = ws.cell(row=r, column=col_idx).value
            try:
                numeric_cols[col_idx] += int(val) if val else 0
            except (TypeError, ValueError):
                pass

    for col_idx, total_val in numeric_cols.items():
        cell = ws.cell(row=current_row, column=col_idx, value=total_val)
        cell.font = total_font
        cell.fill = total_fill
        cell.alignment = cell_align_center
        cell.border = thin_border

    cell = ws.cell(row=current_row, column=9, value="")
    cell.fill = total_fill
    cell.border = thin_border

    wb.save(output_path)
    return output_path


# ==========================================
# تست
# ==========================================

if __name__ == "__main__":
    from database import initialize_database

    initialize_database()

    form_date = "1405/07/24"

    try:
        path = generate_counting_pdf(form_date)
        print(f"✅ PDF ساخته شد: {path}")

        try:
            os.startfile(path)
        except Exception:
            pass
    except Exception as e:
        print(f"❌ خطا: {e}")
