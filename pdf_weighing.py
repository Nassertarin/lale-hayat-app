# -*- coding: utf-8 -*-
"""
تولید PDF و Excel فرم توزین از دیتابیس
"""

import os
import csv
from datetime import datetime
from fpdf import FPDF

import arabic_reshaper
from bidi.algorithm import get_display

try:
    import jdatetime
    HAS_JDATETIME = True
except ImportError:
    HAS_JDATETIME = False


BASE_DIR = os.path.dirname(os.path.abspath(__file__))

FONT_REG_PATH = os.path.join(BASE_DIR, "assets", "fonts", "BNazanin.ttf")
FONT_BOLD_PATH = os.path.join(BASE_DIR, "assets", "fonts", "BNazaninBold.ttf")
LOGO_PATH = os.path.join(BASE_DIR, "assets", "images", "logo.png")

REPORTS_DIR = os.path.join(BASE_DIR, "reports")


def fa(text):
    if not text:
        return ""
    reshaped = arabic_reshaper.reshape(str(text))
    return get_display(reshaped)


def safe_float(value, default=0.0):
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


BLACK = (0, 0, 0)
WHITE = (255, 255, 255)

LEFT_EDGE = 12
RIGHT_EDGE = 198
TOTAL_WIDTH = RIGHT_EDGE - LEFT_EDGE   # 186mm


# ==========================================
# کلاس PDF
# ==========================================

class WeighingFormPDF(FPDF):

    def __init__(self):
        super().__init__(orientation="P", unit="mm", format="A4")
        self.set_auto_page_break(auto=False)
        self.set_margins(12, 12, 12)

        if not os.path.exists(FONT_REG_PATH):
            raise FileNotFoundError(f"فونت پیدا نشد: {FONT_REG_PATH}")

        self.add_font("Nazanin", "", FONT_REG_PATH)

        if os.path.exists(FONT_BOLD_PATH):
            self.add_font("Nazanin", "B", FONT_BOLD_PATH)
            self.has_bold = True
        else:
            self.has_bold = False

    def draw_header(self, form_date, day_name, slaughter_number):
        if os.path.exists(LOGO_PATH):
            try:
                self.image(LOGO_PATH, x=14, y=8, w=18)
            except Exception:
                pass

        self.set_y(10)
        if self.has_bold:
            self.set_font("Nazanin", "B", size=17)
        else:
            self.set_font("Nazanin", size=17)

        self.set_text_color(*BLACK)
        self.cell(186, 11, fa("فرم توزین"), align="C")

        self.set_y(28)
        self.set_font("Nazanin", size=11)
        self.set_text_color(*BLACK)

        date_text = f"تاریخ: {form_date}"
        if day_name:
            date_text += f"  ({day_name})"
        if slaughter_number:
            date_text += f"  —  کشتار #{slaughter_number}"

        self.set_xy(120, 28)
        self.cell(78, 6, fa(date_text), align="R")

        self.set_y(38)

    def draw_section_title(self, title, y):
        if self.has_bold:
            self.set_font("Nazanin", "B", size=12)
        else:
            self.set_font("Nazanin", size=12)

        self.set_text_color(*BLACK)
        self.set_xy(LEFT_EDGE, y)
        self.cell(TOTAL_WIDTH, 7, fa(title), align="R")

    def draw_plucked_table(self, rows, y):
        self.draw_section_title("مرغ پرکنده", y)
        y += 8

        unit = TOTAL_WIDTH / 15

        widths = [
            2.0 * unit,
            3.0 * unit,
            3.0 * unit,
            2.8 * unit,
            2.6 * unit,
            1.6 * unit,
        ]
        headers = [
            "واحد", "وزن خالص", "وزن پر",
            "وزن خالی", "شماره گاری", "ردیف",
        ]

        self.set_font("Nazanin", size=9)
        self.set_text_color(*BLACK)
        self.set_draw_color(*BLACK)

        self.set_xy(LEFT_EDGE, y)
        for h, w in zip(headers, widths):
            self.cell(w, 8.5, fa(h), border=1, align="C")
        self.ln()

        self.set_font("Nazanin", size=9)

        total_empty = 0.0
        total_full = 0.0
        total_net = 0.0

        for row in rows:
            empty_f = safe_float(row["empty_weight"])
            full_f = safe_float(row["full_weight"])
            net_f = safe_float(row["net_weight"])

            self.cell(widths[0], 7, fa("کیلوگرم"),
                      border=1, align="C")
            self.cell(widths[1], 7,
                      f"{net_f:,.2f}" if net_f else "",
                      border=1, align="C")
            self.cell(widths[2], 7,
                      f"{full_f:,.2f}" if full_f else "",
                      border=1, align="C")
            self.cell(widths[3], 7,
                      f"{empty_f:,.2f}" if empty_f else "",
                      border=1, align="C")
            self.cell(widths[4], 7, fa(row["cart_number"]),
                      border=1, align="C")
            self.cell(widths[5], 7, str(row["row_num"]),
                      border=1, align="C")
            self.ln()

            total_empty += empty_f
            total_full += full_f
            total_net += net_f

        self.set_font("Nazanin", size=9)

        self.cell(widths[0], 7.5, fa("کیلوگرم"),
                  border=1, align="C")
        self.cell(widths[1], 7.5, f"{total_net:,.2f}",
                  border=1, align="C")
        self.cell(widths[2], 7.5, f"{total_full:,.2f}",
                  border=1, align="C")
        self.cell(widths[3], 7.5, f"{total_empty:,.2f}",
                  border=1, align="C")
        self.cell(widths[4] + widths[5], 7.5,
                  fa("جمع"), border=1, align="C")

        self.ln(10)
        return self.get_y()

    def draw_alayeshat_table(self, rows, y):
        self.draw_section_title("آلایشات", y)
        y += 8

        unit = TOTAL_WIDTH / 14

        widths = [
            2.0 * unit,
            2.2 * unit,
            2.2 * unit,
            2.2 * unit,
            3.4 * unit,
            2.0 * unit,
        ]
        headers = ["واحد", "وزن خالص", "وزن پر",
                   "وزن خالی", "نام کالا", "ردیف"]

        self.set_font("Nazanin", size=9)
        self.set_text_color(*BLACK)
        self.set_draw_color(*BLACK)

        self.set_xy(LEFT_EDGE, y)
        for h, w in zip(headers, widths):
            self.cell(w, 8.5, fa(h), border=1, align="C")
        self.ln()

        self.set_font("Nazanin", size=9)

        total_empty = 0.0
        total_full = 0.0
        total_net = 0.0

        for row in rows:
            empty_f = safe_float(row["empty_weight"])
            full_f = safe_float(row["full_weight"])
            net_f = safe_float(row["net_weight"])

            self.cell(widths[0], 7, fa("کیلوگرم"),
                      border=1, align="C")
            self.cell(widths[1], 7,
                      f"{net_f:,.2f}" if net_f else "",
                      border=1, align="C")
            self.cell(widths[2], 7,
                      f"{full_f:,.2f}" if full_f else "",
                      border=1, align="C")
            self.cell(widths[3], 7,
                      f"{empty_f:,.2f}" if empty_f else "",
                      border=1, align="C")
            self.cell(widths[4], 7, fa(row["product_name"]),
                      border=1, align="R")
            self.cell(widths[5], 7, str(row["row_num"]),
                      border=1, align="C")
            self.ln()

            total_empty += empty_f
            total_full += full_f
            total_net += net_f

        self.set_font("Nazanin", size=9)

        self.cell(widths[0], 7.5, fa("کیلوگرم"),
                  border=1, align="C")
        self.cell(widths[1], 7.5, f"{total_net:,.2f}",
                  border=1, align="C")
        self.cell(widths[2], 7.5, f"{total_full:,.2f}",
                  border=1, align="C")
        self.cell(widths[3], 7.5, f"{total_empty:,.2f}",
                  border=1, align="C")
        self.cell(widths[4] + widths[5], 7.5,
                  fa("جمع"), border=1, align="C")

        self.ln(10)
        return self.get_y()

    def draw_calc_summary(self, live_weight, total_net_plucked,
                          loss_percent, y):
        """
        بخش محاسبه — نسخه‌ی نهایی:
        - فقط دو خط فرمول (عمومی + جایگذاری)
        - استفاده از براکت [ ] به جای پرانتز بیرونی
        - خط نهایی درصد افت
        """
        self.draw_section_title("محاسبه", y)
        y += 9

        self.set_font("Nazanin", size=10)
        self.set_text_color(*BLACK)

        # ── خط ۱: فرمول عمومی ──
        self.set_xy(LEFT_EDGE, y)
        self.cell(
            TOTAL_WIDTH, 6.5,
            fa("۱۰۰ × [ (وزن خالص پرکنده ÷ وزن زنده باسکول) - ۱ ] = درصد افت"),
            align="R",
        )
        y += 8.0

        # ── خط ۲: جایگذاری اعداد ──
        self.set_xy(LEFT_EDGE, y)
        self.cell(
            TOTAL_WIDTH, 6.5,
            fa(
                f"۱۰۰ × [ ({total_net_plucked:,.2f} ÷ {live_weight:,.2f}) - ۱ ] = درصد افت"
            ),
            align="R",
        )
        y += 10.0

        # ── خط ۳: درصد افت نهایی ──
        if self.has_bold:
            self.set_font("Nazanin", "B", size=12)
        else:
            self.set_font("Nazanin", size=12)

        self.set_xy(LEFT_EDGE, y)
        self.cell(
            TOTAL_WIDTH, 8,
            fa(f"{loss_percent:.2f}% = درصد افت"),
            align="R",
        )
        y += 8

        return y


# ==========================================
# تولید PDF
# ==========================================

def generate_weighing_pdf(slaughter_id, output_path=None):
    from functions import (
        get_weighing_form,
        get_slaughter,
    )

    form = get_weighing_form(slaughter_id)

    if not form:
        raise ValueError(f"فرم توزین برای کشتار #{slaughter_id} یافت نشد.")

    slaughter = get_slaughter(slaughter_id)
    if not slaughter:
        raise ValueError(f"کشتار #{slaughter_id} یافت نشد.")

    form_date = slaughter.get("form_date", "")
    day_name = slaughter.get("day_name", "")
    slaughter_number = slaughter.get("slaughter_number", "")

    pdf = WeighingFormPDF()
    pdf.add_page()

    pdf.draw_header(form_date, day_name, slaughter_number)

    y = pdf.get_y() + 6

    y = pdf.draw_plucked_table(form["plucked_rows"], y)

    y += 4
    y = pdf.draw_alayeshat_table(form["alayeshat_rows"], y)

    y += 4

    # ── محاسبه وزن خالص فقط مرغ پرکنده ──
    total_net_plucked = sum(
        safe_float(r["net_weight"]) for r in form["plucked_rows"]
    )

    live_weight = safe_float(form.get("live_weight", 0))

    y = pdf.draw_calc_summary(
        live_weight=live_weight,
        total_net_plucked=total_net_plucked,
        loss_percent=form["loss_percent"],
        y=y,
    )

    if output_path is None:
        os.makedirs(REPORTS_DIR, exist_ok=True)
        safe_date = form_date.replace("/", "-")
        ts = datetime.now().strftime("%H%M%S")
        filename = f"weighing_{safe_date}_K{slaughter_number}_{ts}.pdf"
        output_path = os.path.join(REPORTS_DIR, filename)

    pdf.output(output_path)
    return output_path


# ==========================================
# تولید Excel (xlsx) — با openpyxl
# ==========================================

def generate_weighing_excel(slaughter_id, output_path=None):
    """
    تولید فایل Excel (xlsx) از فرم ۲ (توزین) با استایل‌دهی مناسب (مشابه قبض باسکول).
    """
    from functions import get_weighing_form, get_slaughter

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ImportError(
            "کتابخانه openpyxl نصب نیست. لطفاً با دستور 'pip install openpyxl' نصب کنید."
        )

    form = get_weighing_form(slaughter_id)
    if not form:
        raise ValueError(f"فرم توزین برای کشتار #{slaughter_id} یافت نشد.")

    slaughter = get_slaughter(slaughter_id)
    if not slaughter:
        raise ValueError(f"کشتار #{slaughter_id} یافت نشد.")

    form_date = slaughter.get("form_date", "")
    day_name = slaughter.get("day_name", "")
    slaughter_number = slaughter.get("slaughter_number", "")

    if output_path is None:
        os.makedirs(REPORTS_DIR, exist_ok=True)
        safe_date = form_date.replace("/", "-")
        ts = datetime.now().strftime("%H%M%S")
        filename = f"weighing_{safe_date}_K{slaughter_number}_{ts}.xlsx"
        output_path = os.path.join(REPORTS_DIR, filename)

    # ── استایل‌ها (مشابه قبض باسکول) ──
    font_title = Font(name="B Nazanin", size=16, bold=True, color="147A6F")
    font_subtitle = Font(name="B Nazanin", size=11, color="333333")
    font_section = Font(name="B Nazanin", size=13, bold=True, color="147A6F")
    font_header = Font(name="B Nazanin", size=11, bold=True, color="FFFFFF")
    font_normal = Font(name="B Nazanin", size=11, color="000000")
    font_bold = Font(name="B Nazanin", size=11, bold=True, color="000000")
    font_calc = Font(name="B Nazanin", size=11, color="000000")
    font_calc_result = Font(name="B Nazanin", size=12, bold=True, color="147A6F")

    fill_primary = PatternFill(start_color="147A6F", end_color="147A6F", fill_type="solid")
    fill_total = PatternFill(start_color="E0F2F1", end_color="E0F2F1", fill_type="solid")

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    border_thin = Border(
        left=Side(style="thin", color="999999"),
        right=Side(style="thin", color="999999"),
        top=Side(style="thin", color="999999"),
        bottom=Side(style="thin", color="999999"),
    )
    border_separator = Border(
        bottom=Side(style="medium", color="147A6F"),
    )

    # ── ساخت Workbook ──
    wb = Workbook()
    ws = wb.active
    ws.title = "توزین"
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False

    # ── تنظیم عرض ستون‌ها ──
    col_widths = {
        "A": 10,   # ردیف / برچسب
        "B": 18,   # شماره گاری / نام کالا
        "C": 16,   # وزن خالی
        "D": 16,   # وزن پر
        "E": 16,   # وزن خالص
        "F": 14,   # واحد
    }
    for col, w in col_widths.items():
        ws.column_dimensions[col].width = w

    # ══════════════════════════════════════
    # عنوان
    # ══════════════════════════════════════
    ws.merge_cells("A1:F1")
    c = ws["A1"]
    c.value = "فرم توزین"
    c.font = font_title
    c.alignment = align_center
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:F2")
    c = ws["A2"]
    date_text = f"تاریخ: {form_date}"
    if day_name:
        date_text += f"  ({day_name})"
    if slaughter_number:
        date_text += f"  —  کشتار #{slaughter_number}"
    c.value = date_text
    c.font = font_subtitle
    c.alignment = align_center
    ws.row_dimensions[2].height = 22

    # خط جداکننده‌ی زیر عنوان (مشابه قبض)
    ws.merge_cells("A3:F3")
    c = ws["A3"]
    c.border = border_separator

    current_row = 5

    # ══════════════════════════════════════
    # جدول ۱: مرغ پرکنده
    # ══════════════════════════════════════
    ws.merge_cells(f"A{current_row}:F{current_row}")
    c = ws[f"A{current_row}"]
    c.value = "مرغ پرکنده"
    c.font = font_section
    c.alignment = align_right
    current_row += 1

    headers1 = ["ردیف", "شماره گاری", "وزن خالی", "وزن پر", "وزن خالص", "واحد"]

    for col_idx, header in enumerate(headers1, start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=header)
        cell.font = font_header
        cell.fill = fill_primary
        cell.alignment = align_center
        cell.border = border_thin
    ws.row_dimensions[current_row].height = 24
    current_row += 1

    total_empty = 0.0
    total_full = 0.0
    total_net = 0.0

    for row in form["plucked_rows"]:
        empty_f = safe_float(row.get("empty_weight", 0))
        full_f = safe_float(row.get("full_weight", 0))
        net_f = safe_float(row.get("net_weight", 0))

        total_empty += empty_f
        total_full += full_f
        total_net += net_f

        values = [
            row["row_num"],
            row.get("cart_number", ""),
            empty_f,
            full_f,
            net_f,
            "کیلوگرم",
        ]

        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=value)
            cell.font = font_normal
            # ستون شماره گاری راست‌چین، بقیه وسط‌چین
            cell.alignment = align_right if col_idx == 2 else align_center
            cell.border = border_thin
        current_row += 1

    # ردیف جمع
    values_total = ["جمع", "", total_empty, total_full, total_net, "کیلوگرم"]
    for col_idx, value in enumerate(values_total, start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=value)
        cell.font = font_bold
        cell.fill = fill_total
        cell.alignment = align_center
        cell.border = border_thin
    current_row += 2

    # ══════════════════════════════════════
    # جدول ۲: آلایشات
    # ══════════════════════════════════════
    ws.merge_cells(f"A{current_row}:F{current_row}")
    c = ws[f"A{current_row}"]
    c.value = "آلایشات"
    c.font = font_section
    c.alignment = align_right
    current_row += 1

    headers2 = ["ردیف", "نام کالا", "وزن خالی", "وزن پر", "وزن خالص", "واحد"]

    for col_idx, header in enumerate(headers2, start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=header)
        cell.font = font_header
        cell.fill = fill_primary
        cell.alignment = align_center
        cell.border = border_thin
    ws.row_dimensions[current_row].height = 24
    current_row += 1

    total_empty = 0.0
    total_full = 0.0
    total_net = 0.0

    for row in form["alayeshat_rows"]:
        empty_f = safe_float(row.get("empty_weight", 0))
        full_f = safe_float(row.get("full_weight", 0))
        net_f = safe_float(row.get("net_weight", 0))

        total_empty += empty_f
        total_full += full_f
        total_net += net_f

        values = [
            row["row_num"],
            row.get("product_name", ""),
            empty_f,
            full_f,
            net_f,
            "کیلوگرم",
        ]

        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=value)
            cell.font = font_normal
            cell.alignment = align_right if col_idx == 2 else align_center
            cell.border = border_thin
        current_row += 1

    # ردیف جمع
    values_total = ["جمع", "", total_empty, total_full, total_net, "کیلوگرم"]
    for col_idx, value in enumerate(values_total, start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=value)
        cell.font = font_bold
        cell.fill = fill_total
        cell.alignment = align_center
        cell.border = border_thin
    current_row += 2

    # ══════════════════════════════════════
    # بخش محاسبه (با استایل کادر مشابه توضیحات قبض)
    # ══════════════════════════════════════
    ws.merge_cells(f"A{current_row}:F{current_row}")
    c = ws[f"A{current_row}"]
    c.value = "محاسبه"
    c.font = font_section
    c.alignment = align_right
    current_row += 1

    # ── محاسبه وزن خالص فقط مرغ پرکنده ──
    total_net_plucked = sum(
        safe_float(r.get("net_weight", 0)) for r in form["plucked_rows"]
    )
    live_weight = safe_float(form.get("live_weight", 0))
    loss_percent = safe_float(form.get("loss_percent", 0))

    calc_rows = [
        (
            "فرمول",
            "۱۰۰ × [ (وزن خالص پرکنده ÷ وزن زنده باسکول) - ۱ ] = درصد افت",
            font_calc,
        ),
        (
            "جایگذاری",
            f"۱۰۰ × [ ({total_net_plucked:,.2f} ÷ {live_weight:,.2f}) - ۱ ] = درصد افت",
            font_calc,
        ),
        (
            "درصد افت",
            f"{loss_percent:.2f}%",
            font_calc_result,
        ),
    ]

    for label, value, value_font in calc_rows:
        # ستون برچسب (A:B)
        ws.merge_cells(
            start_row=current_row, start_column=1,
            end_row=current_row, end_column=2
        )
        c = ws.cell(row=current_row, column=1, value=label)
        c.font = font_bold
        c.alignment = align_right
        c.border = border_thin

        # ستون مقدار (C:F)
        ws.merge_cells(
            start_row=current_row, start_column=3,
            end_row=current_row, end_column=6
        )
        c = ws.cell(row=current_row, column=3, value=value)
        c.font = value_font
        c.alignment = align_right
        c.border = border_thin

        # اعمال حاشیه روی همه‌ی سلول‌های ادغام‌شده
        for col_idx in range(1, 7):
            ws.cell(row=current_row, column=col_idx).border = border_thin

        ws.row_dimensions[current_row].height = 26
        current_row += 1

    # ── ذخیره ──
    wb.save(output_path)
    return output_path


# ==========================================
# خروجی CSV (سازگاری با کد قبلی)
# ==========================================

def generate_weighing_csv(slaughter_id, output_path=None):
    """
    تولید فایل CSV از فرم ۲ (توزین).
    """
    from functions import get_weighing_form, get_slaughter

    form = get_weighing_form(slaughter_id)
    if not form:
        raise ValueError(f"فرم توزین یافت نشد.")

    slaughter = get_slaughter(slaughter_id)
    if not slaughter:
        raise ValueError(f"کشتار یافت نشد.")

    if output_path is None:
        os.makedirs(REPORTS_DIR, exist_ok=True)
        safe_date = slaughter["form_date"].replace("/", "-")
        ts = datetime.now().strftime("%H%M%S")
        filename = f"weighing_{safe_date}_K{slaughter['slaughter_number']}_{ts}.csv"
        output_path = os.path.join(REPORTS_DIR, filename)

    # ── محاسبه وزن خالص فقط مرغ پرکنده ──
    total_net_plucked = sum(
        safe_float(r.get("net_weight", 0)) for r in form["plucked_rows"]
    )
    live_weight = safe_float(form.get("live_weight", 0))

    with open(output_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)

        writer.writerow(["فرم توزین"])
        writer.writerow([
            f"کشتار #{slaughter['slaughter_number']}",
            f"{slaughter['day_name']} {slaughter['form_date']}",
        ])
        writer.writerow([])

        writer.writerow(["مرغ پرکنده"])
        writer.writerow([
            "ردیف", "شماره گاری", "وزن خالی", "وزن پر", "خالص"
        ])

        for i, r in enumerate(form["plucked_rows"], start=1):
            writer.writerow([
                i,
                r["cart_number"],
                f"{r['empty_weight']:.2f}",
                f"{r['full_weight']:.2f}",
                f"{r['net_weight']:.2f}",
            ])

        writer.writerow([])

        writer.writerow(["آلایشات"])
        writer.writerow([
            "ردیف", "محصول", "وزن خالی", "وزن پر", "خالص"
        ])

        for i, r in enumerate(form["alayeshat_rows"], start=1):
            writer.writerow([
                i,
                r["product_name"],
                f"{r['empty_weight']:.2f}",
                f"{r['full_weight']:.2f}",
                f"{r['net_weight']:.2f}",
            ])

        writer.writerow([])

        writer.writerow(["محاسبه"])
        writer.writerow([
            "فرمول",
            "۱۰۰ × [ (وزن خالص پرکنده ÷ وزن زنده باسکول) - ۱ ] = درصد افت",
        ])
        writer.writerow([
            "جایگذاری",
            f"۱۰۰ × [ ({total_net_plucked:,.2f} ÷ {live_weight:,.2f}) - ۱ ] = درصد افت",
        ])
        writer.writerow([
            "درصد افت",
            f"{form['loss_percent']:.2f}%",
        ])

    return output_path


# ==========================================
# تست
# ==========================================

if __name__ == "__main__":
    from database import initialize_database

    initialize_database()

    slaughter_id = 1

    try:
        path = generate_weighing_pdf(slaughter_id)
        print(f"✅ PDF ساخته شد: {path}")

        try:
            os.startfile(path)
        except Exception:
            pass
    except Exception as e:
        print(f"❌ خطا: {e}")
